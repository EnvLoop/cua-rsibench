"""Public-library generation and original independent controls, no Office UI."""
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from openpyxl import load_workbook
from ppt_wdi_factory import verify as ppt_verify
from ppt_wdi_factory.build_portable_v1 import build as build_ppt
from sec_excel_factory.build_portable_integrated_v1 import generate,formulas
from sec_excel_factory.prepare_integrated_cases import allocate_cases,packages
from sec_excel_factory.verify_integrated_candidate import verify as sec_verify,target_addresses


class PortableOfficeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)

    def test_portable_sec_original_oracle_positive_wrong_and_collateral(self):
        all_cases=allocate_cases(packages());cases=[all_cases[i] for i in (0,20,40)]
        output=self.root/'sec';generate(cases,output)
        for case in cases:
            folder=output/case['split']/case['case_id'];actor=folder/'actor.xlsx';reference=folder/'reference.xlsx'
            self.assertEqual({tuple(key.split('!')) for key in formulas(case)},target_addresses(case['split']))
            self.assertFalse(sec_verify(actor,actor,case)['pass'])
            self.assertTrue(sec_verify(reference,actor,case)['pass'])
            wrong=folder/'wrong.xlsx';book=load_workbook(reference);book['Review']['B5']='=1';book.save(wrong)
            self.assertFalse(sec_verify(wrong,actor,case)['pass'])
            collateral=folder/'collateral.xlsx';book=load_workbook(reference);book['Annual 10-K']['A5']='Unauthorized source rewrite';book.save(collateral)
            self.assertFalse(sec_verify(collateral,actor,case)['pass'])

    def test_portable_ppt_original_source_chart_and_triad_controls(self):
        source=Path(__file__).resolve().parents[1]/'work/full-study/office-ppt-v2-native-qualification-root-20261001.private/neutral-lifecycle-v4-prepared.private/package.private'
        if not (source/'task.private.json').is_file():self.skipTest('Actual private WDI source fixture absent')
        package=self.root/'ppt';package.mkdir();task=json.loads((source/'task.private.json').read_bytes())
        for name in ('task.private.json','source-snapshot.private.json','source-provenance.private.json','source-country.private.zip'):
            if (source/name).is_file():shutil.copyfile(source/name,package/name)
        build_ppt(task,package/'source.pptx');result=ppt_verify.calibrate(package)
        self.assertTrue(result['offline_controls_pass'])
        self.assertNotEqual((source/'source.pptx').read_bytes(),(package/'source.pptx').read_bytes())


if __name__=='__main__':unittest.main()
