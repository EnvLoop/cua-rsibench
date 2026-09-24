"""Build an intentionally faulty, source-grounded Excel model for blind repair.

This public fixture is a development example, not an official hidden task. All
filing facts come from the already frozen SEC companyfacts snapshots.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from extract_sec import EXTRA_METRICS, ISSUERS, METRICS, read_source


HERE = Path(__file__).resolve().parent
PERIODS = {
    "Apple": {"start": "2023-10-01", "nine_end": "2024-06-29", "nine_accession": "0000320193-24-000081"},
    "Microsoft": {"start": "2023-07-01", "nine_end": "2024-03-31", "nine_accession": "0000950170-24-048288"},
}
METRIC_COLUMNS = {"Revenue": "D", "Operating income": "E", "Operating cash flow": "F", "Capital expenditures": "G", "Assets": "H", "Liabilities": "I", "Equity": "J"}
BRIDGE_METRICS = ("Revenue", "Operating cash flow", "Capital expenditures", "Operating income")
SHEETS = ("Board Review", "Historical", "Q4 Bridge", "Drivers", "Forecast", "Raw SEC")


def records() -> list[dict]:
    rows = []
    for issuer, spec in ISSUERS.items():
        data = read_source(spec)
        endings = set(spec["periods"].values()) | {PERIODS[issuer]["nine_end"]}
        for metric, concept in METRICS.items():
            for fact in data["facts"]["us-gaap"][concept]["units"]["USD"]:
                if fact.get("end") not in endings or fact.get("form") not in ("10-K", "10-K/A", "10-Q", "10-Q/A"):
                    continue
                if fact.get("filed", "") > "2025-12-31":
                    continue
                rows.append(_record(issuer, spec, metric, concept, "USD", fact))
        # Genuine per-share and share-count entries make unit matching necessary.
        for metric, (concept, unit) in EXTRA_METRICS.items():
            for fact in data["facts"]["us-gaap"][concept]["units"][unit]:
                if fact.get("end") in endings and fact.get("form") in ("10-K", "10-Q") and fact.get("filed", "") <= "2025-12-31":
                    rows.append(_record(issuer, spec, metric, concept, unit, fact))
    unique = {r["id"]: r for r in rows}
    if len(unique) != len(rows):
        raise ValueError("source_record_id_collision_or_duplicate")
    return sorted(rows, key=lambda r: (r["issuer"], r["metric"], r["end"], r["start"], r["filed"], r["id"]))


def _record(issuer: str, spec: dict, metric: str, concept: str, unit: str, fact: dict) -> dict:
    payload = {k: fact.get(k) for k in ("start", "end", "val", "accn", "form", "filed", "frame")}
    rid = hashlib.sha256(f"{spec['cik']}:{concept}:{unit}:{json.dumps(payload, sort_keys=True)}".encode()).hexdigest()[:20]
    return {"id": rid, "issuer": issuer, "cik": spec["cik"], "metric": metric, "concept": concept,
            "start": fact.get("start", ""), "end": fact["end"], "filed": fact["filed"],
            "form": fact["form"], "accession": fact["accn"], "fy": fact.get("fy", ""),
            "fp": fact.get("fp", ""), "frame": fact.get("frame", ""), "unit": unit,
            "value": fact["val"]}


def _find(rows: list[dict], issuer: str, metric: str, *, start: str | None, end: str, accession: str) -> int:
    matches = [i + 5 for i, r in enumerate(rows) if r["issuer"] == issuer and r["metric"] == metric
               and r["start"] == (start or "") and r["end"] == end and r["accession"] == accession
               and r["unit"] == "USD"]
    if len(matches) != 1:
        raise ValueError(f"source_context_not_unique:{issuer}:{metric}:{start}:{end}:{accession}:{matches}")
    return matches[0]


def formula_plan(rows: list[dict]) -> tuple[dict[tuple[str, str], str], dict[tuple[str, str], str]]:
    formulas: dict[tuple[str, str], str] = {}
    defects: dict[tuple[str, str], str] = {}

    def put(sheet: str, cell: str, formula: str) -> None:
        formulas[(sheet, cell)] = formula

    hist = {("Apple", "2023"): 5, ("Apple", "2024"): 6, ("Microsoft", "2023"): 7, ("Microsoft", "2024"): 8}
    for (issuer, year), row in hist.items():
        spec = ISSUERS[issuer]
        fy_start = "2022-09-25" if (issuer, year) == ("Apple", "2023") else (
            "2023-10-01" if issuer == "Apple" else ("2022-07-01" if year == "2023" else "2023-07-01"))
        for metric, col in METRIC_COLUMNS.items():
            source = _find(rows, issuer, metric, start=None if metric in ("Assets", "Liabilities", "Equity") else fy_start,
                           end=spec["periods"][year], accession=spec["filing"])
            put("Historical", f"{col}{row}", f"='Raw SEC'!O{source}/1000000")
        put("Historical", f"K{row}", f"=H{row}-I{row}-J{row}")

    bridge_rows = {}
    for issuer, first in (("Apple", 5), ("Microsoft", 9)):
        period = PERIODS[issuer]
        for offset, metric in enumerate(BRIDGE_METRICS):
            row = first + offset
            bridge_rows[(issuer, metric)] = row
            ytd = _find(rows, issuer, metric, start=period["start"], end=period["nine_end"], accession=period["nine_accession"])
            put("Q4 Bridge", f"C{row}", f"='Raw SEC'!O{ytd}/1000000")
            put("Q4 Bridge", f"D{row}", f"=Historical!{METRIC_COLUMNS[metric]}{hist[(issuer, '2024')]}")
            put("Q4 Bridge", f"E{row}", f"=D{row}-C{row}")
        fcf_row = 14 if issuer == "Apple" else 15
        ocf, capex = bridge_rows[(issuer, "Operating cash flow")], bridge_rows[(issuer, "Capital expenditures")]
        for col in "CDE":
            put("Q4 Bridge", f"{col}{fcf_row}", f"={col}{ocf}-{col}{capex}")

    for issuer, driver, y25, y26, board in (("Apple", 5, 5, 6, 5), ("Microsoft", 6, 7, 8, 6)):
        past, current = hist[(issuer, "2023")], hist[(issuer, "2024")]
        put("Drivers", f"B{driver}", f"=Historical!D{current}/Historical!D{past}-1")
        put("Drivers", f"C{driver}", f"=Historical!F{current}/Historical!D{current}")
        put("Drivers", f"D{driver}", f"=Historical!G{current}/Historical!D{current}")
        put("Drivers", f"E{driver}", f"=B{driver}/2")
        for row, prior, growth in ((y25, f"Historical!D{current}", "B"), (y26, f"D{y25}", "E")):
            put("Forecast", f"C{row}", f"={prior}")
            put("Forecast", f"D{row}", f"=C{row}*(1+Drivers!{growth}{driver})")
            put("Forecast", f"E{row}", f"=D{row}*Drivers!C{driver}")
            put("Forecast", f"F{row}", f"=D{row}*Drivers!D{driver}*Drivers!F{driver}")
            put("Forecast", f"G{row}", f"=E{row}-F{row}")
            put("Forecast", f"H{row}", f"=D{row}*Historical!E{current}/Historical!D{current}")
            put("Forecast", f"I{row}", f"=G{row}/D{row}")
        q4rev, q4fcf = bridge_rows[(issuer, "Revenue")], (14 if issuer == "Apple" else 15)
        for col, expression in {
            "B": f"Historical!D{current}", "C": f"Historical!F{current}-Historical!G{current}",
            "D": f"'Q4 Bridge'!E{q4rev}", "E": f"'Q4 Bridge'!E{q4fcf}",
            "F": f"E{board}/D{board}", "G": f"Forecast!D{y26}", "H": f"Forecast!G{y26}",
            "I": f"Historical!K{current}",
        }.items():
            put("Board Review", f"{col}{board}", f"={expression}")

    # A same-value later comparative is invisible to a static value-only check.
    apple_annual = _find(rows, "Apple", "Revenue", start="2023-10-01", end="2024-09-28", accession=ISSUERS["Apple"]["filing"])
    alternatives = [i + 5 for i, r in enumerate(rows) if r["issuer"] == "Apple" and r["metric"] == "Revenue"
                    and r["start"] == "2023-10-01" and r["end"] == "2024-09-28"
                    and r["accession"] != ISSUERS["Apple"]["filing"] and r["value"] == rows[apple_annual - 5]["value"]]
    if not alternatives:
        raise ValueError("same_value_comparative_missing")
    defects[("Historical", "D6")] = f"='Raw SEC'!O{alternatives[0]}/1000000"
    defects[("Historical", "F7")] = "='Raw SEC'!O" + str(_find(rows, "Microsoft", "Operating cash flow", start="2023-07-01", end="2024-06-30", accession=ISSUERS["Microsoft"]["filing"])) + "/1000000"
    apple_q3_quarter = [i + 5 for i, r in enumerate(rows) if r["issuer"] == "Apple" and r["metric"] == "Capital expenditures"
                        and r["end"] == PERIODS["Apple"]["nine_end"] and r["accession"] == PERIODS["Apple"]["nine_accession"]
                        and r["start"] != PERIODS["Apple"]["start"]]
    # Capex companyfacts often has YTD only; use a genuine wrong accession if no quarter-only fact exists.
    wrong_capex = apple_q3_quarter[0] if apple_q3_quarter else _find(rows, "Apple", "Capital expenditures", start="2023-10-01", end="2024-09-28", accession=ISSUERS["Apple"]["filing"])
    defects[("Q4 Bridge", "C7")] = f"='Raw SEC'!O{wrong_capex}/1000000"
    defects[("Q4 Bridge", "E11")] = "=D11+C11"
    defects[("Drivers", "D5")] = "=Historical!G6/Historical!D5"
    defects[("Forecast", "E5")] = "=D5*Drivers!C6"
    defects[("Forecast", "D8")] = "=C8*(1+Drivers!B6)"
    defects[("Board Review", "F5")] = "=C5/B5"
    defects[("Board Review", "H6")] = "=Forecast!G7"
    if any(formulas[key] == bad for key, bad in defects.items()):
        raise ValueError("fault_has_no_effect_on_formula")
    return formulas, defects


def _heading(ws, title: str, description: str, headers: list[str]) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos", size=15, bold=True, color="173A5E")
    ws["A2"] = description
    ws["A2"].font = Font(name="Aptos", size=10, color="5A6A78")
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value)
        cell.fill = PatternFill("solid", fgColor="173A5E")
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 34
    ws.column_dimensions["A"].width = 22
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 21
    ws.freeze_panes = "B5"


def workbook(rows: list[dict], formulas: dict, defects: dict, *, faulty: bool) -> Workbook:
    wb = Workbook()
    wb.properties.created = datetime(2024, 12, 31, 0, 0, 0)
    wb.properties.modified = datetime(2024, 12, 31, 0, 0, 0)
    wb.active.title = SHEETS[0]
    for name in SHEETS[1:]:
        wb.create_sheet(name)
    board, historical, bridge, drivers, forecast, raw = [wb[name] for name in SHEETS]
    _heading(board, "FY2024 board review / FY2026 scenario", "USD millions; issuer fiscal year ends differ; Q4 is implied from 10-K less nine-month 10-Q.",
             ["Issuer", "FY24 revenue", "FY24 FCF", "Implied Q4 revenue", "Implied Q4 FCF", "Q4 FCF margin", "FY26 revenue", "FY26 FCF", "Balance check"])
    for row, issuer in ((5, "Apple"), (6, "Microsoft")):
        board[f"A{row}"] = issuer
    _heading(historical, "Historical filing selection", "Use FY2024 10-K comparative facts for FY2023 and FY2024; USD millions.",
             ["Issuer", "Fiscal year", "Period end", *METRIC_COLUMNS, "Balance check"])
    for row, issuer, year, end in ((5, "Apple", 2023, "2023-09-30"), (6, "Apple", 2024, "2024-09-28"),
                                    (7, "Microsoft", 2023, "2023-06-30"), (8, "Microsoft", 2024, "2024-06-30")):
        historical[f"A{row}"] = issuer
        historical[f"B{row}"] = year
        historical[f"C{row}"] = end
    _heading(bridge, "Implied fiscal Q4 bridge", "Nine-month YTD from the original FY2024 Q3 10-Q; full year from the FY2024 10-K.",
             ["Issuer", "Metric", "Nine-month YTD", "Fiscal full year", "Implied Q4"])
    for issuer, first in (("Apple", 5), ("Microsoft", 9)):
        for offset, metric in enumerate(BRIDGE_METRICS):
            bridge[f"A{first + offset}"] = issuer
            bridge[f"B{first + offset}"] = metric
    for row, issuer in ((14, "Apple"), (15, "Microsoft")):
        bridge[f"A{row}"] = issuer
        bridge[f"B{row}"] = "Free cash flow"
    _heading(drivers, "Forecast drivers and scenario", "FY25 growth repeats FY24 growth; FY26 growth halves it. Capex scenario is a modeled input.",
             ["Issuer", "FY25 revenue growth", "FY24 OCF/revenue", "FY24 capex/revenue", "FY26 revenue growth", "Capex multiplier"])
    for row, issuer, multiplier in ((5, "Apple", 1.173), (6, "Microsoft", 1.0)):
        drivers[f"A{row}"] = issuer
        drivers[f"F{row}"] = multiplier
        drivers[f"F{row}"].fill = PatternFill("solid", fgColor="FFF0BE")
    _heading(forecast, "FY2025–2026 operating forecast", "Illustrative, not a company forecast; USD millions.",
             ["Issuer", "Fiscal year", "Prior revenue", "Revenue", "OCF", "Capex", "FCF", "Operating income", "FCF margin"])
    for row, issuer, year in ((5, "Apple", 2025), (6, "Apple", 2026), (7, "Microsoft", 2025), (8, "Microsoft", 2026)):
        forecast[f"A{row}"] = issuer
        forecast[f"B{row}"] = year
    _heading(raw, "Frozen SEC companyfacts records", "All records are copied from the pinned public SEC snapshots; select by issuer, unit, period, form and accession.",
             ["Record ID", "Issuer", "CIK", "Metric", "Concept", "Period start", "Period end", "Filed", "Form", "Accession", "FY", "FP", "Frame", "Unit", "Reported value"])
    raw.column_dimensions["D"].width = 25
    raw.column_dimensions["E"].width = 46
    raw.column_dimensions["J"].width = 28
    raw.column_dimensions["O"].width = 24
    for row_num, r in enumerate(rows, 5):
        vals = [r[k] for k in ("id", "issuer", "cik", "metric", "concept", "start", "end", "filed", "form", "accession", "fy", "fp", "frame", "unit", "value")]
        for col, value in enumerate(vals, 1):
            raw.cell(row_num, col, value)
        raw[f"C{row_num}"].number_format = "0000000000"
        raw[f"O{row_num}"].number_format = "#,##0;(#,##0);-"
    table = Table(displayName="SecFactsBridge", ref=f"A4:O{len(rows) + 4}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    raw.add_table(table)
    for (sheet, addr), formula in formulas.items():
        cell = wb[sheet][addr]
        cell.value = defects.get((sheet, addr), formula) if faulty else formula
        cell.font = Font(name="Aptos", size=10, color="202A33")
        cell.number_format = "0.0%" if (sheet, addr[0]) in (("Board Review", "F"), ("Forecast", "I")) or (sheet == "Drivers") else "#,##0.0;(#,##0.0);-"
    for cell in board["I"][4:6]:
        cell.number_format = "0.00;[Red](0.00);-"
    return wb


def _save_deterministic(wb: Workbook, path: Path) -> None:
    wb.save(path)
    data = io.BytesIO()
    with ZipFile(path) as original, ZipFile(data, "w") as normalized:
        for info in original.infolist():
            info.date_time = (2024, 12, 31, 0, 0, 0)
            payload = original.read(info.filename)
            if info.filename == "docProps/core.xml":
                payload, replacements = re.subn(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2024-12-31T00:00:00Z\g<2>", payload,
                )
                if replacements != 1:
                    raise ValueError("unexpected_core_modified_metadata")
            normalized.writestr(info, payload)
    path.write_bytes(data.getvalue())


def build(out: Path) -> dict:
    source_rows = records()
    formulas, defects = formula_plan(source_rows)
    actor, private = out / "actor", out / "private"
    actor.mkdir(parents=True, exist_ok=True)
    private.mkdir(parents=True, exist_ok=True)
    _save_deterministic(workbook(source_rows, formulas, defects, faulty=True), actor / "task.xlsx")
    _save_deterministic(workbook(source_rows, formulas, defects, faulty=False), private / "reference.xlsx")
    (actor / "task.md").write_text((HERE / "BRIDGE_AUDIT_TASK.md").read_text())
    private_manifest = {"schema": "sec-bridge-blind-audit-v1", "injected_faults": [f"{sheet}!{addr}" for sheet, addr in defects],
                        "source_rows": len(source_rows), "formula_cells": len(formulas),
                        "source_sha256": {name: spec["sha256"] for name, spec in ISSUERS.items()}}
    (private / "fault_manifest.json").write_text(json.dumps(private_manifest, indent=2) + "\n")
    return {"source_rows": len(source_rows), "formula_cells": len(formulas), "injected_faults": len(defects), "output": str(out)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.output), indent=2))
