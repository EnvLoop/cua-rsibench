"""Build a blind retail inventory/payables audit from frozen SEC records.

This is a public development fixture. An official hidden final would require
unpublished fault locations, disjoint issuer cohorts and real GUI admission.
"""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from build_bridge_audit import _save_deterministic
from extract_retail_sec import HERE, ISSUERS


SOURCE = HERE / "sources/retail_excerpt.json"


def records() -> list[dict]:
    data = json.loads(SOURCE.read_text())
    if data["schema"] != "sec-retail-working-capital-source-v1" or len(data["source_meta"]) != 2:
        raise ValueError("retail_source_identity_changed")
    rows = data["records"]
    if len(rows) != len({r["id"] for r in rows}):
        raise ValueError("duplicate_source_record_id")
    return rows


def _find(rows: list[dict], ticker: str, metric: str, year: str, *, prior_accession: bool = False) -> int:
    spec = ISSUERS[ticker]
    end = spec["ends"][year]
    accession = spec["filing_2023"] if year == "2022" or prior_accession else spec["filing_2024"]
    matches = [i for i, r in enumerate(rows, 5) if (r["issuer"], r["metric"], r["end"], r["accession"], r["form"], r["unit"])
               == (ticker, metric, end, accession, "10-K", "USD")]
    if len(matches) != 1:
        raise ValueError(f"raw_reference_ambiguity:{ticker}:{metric}:{year}:{len(matches)}")
    return matches[0]


def plan(rows: list[dict]) -> tuple[dict[tuple[str, str], str], dict[tuple[str, str], str], dict[tuple[str, str], int]]:
    formulas: dict[tuple[str, str], str] = {}
    faults: dict[tuple[str, str], str] = {}
    days: dict[tuple[str, str], int] = {}
    for ticker, year, row in (("COST", "2023", 5), ("COST", "2024", 6),
                              ("WMT", "2023", 7), ("WMT", "2024", 8)):
        current = _find(rows, ticker, "Revenue", year)
        start = next(r["start"] for i, r in enumerate(rows, 5) if i == current)
        end = ISSUERS[ticker]["ends"][year]
        days[("History", f"C{row}")] = (date.fromisoformat(end) - date.fromisoformat(start)).days + 1
        previous = str(int(year) - 1)
        for col, metric, fact_year in (("D", "Revenue", year), ("E", "Cost of sales", year),
                                       ("F", "Inventory", previous), ("G", "Inventory", year),
                                       ("H", "Trade payables", previous), ("I", "Trade payables", year)):
            formulas[("History", f"{col}{row}")] = f"='Raw SEC'!N{_find(rows, ticker, metric, fact_year)}/1000000"
        for col, formula in {
            "J": f"=(F{row}+G{row})/2", "K": f"=(H{row}+I{row})/2",
            "L": f"=J{row}/E{row}*C{row}", "M": f"=K{row}/E{row}*C{row}",
            "N": f"=G{row}-I{row}", "O": f"=(D{row}-E{row})/D{row}",
        }.items():
            formulas[("History", f"{col}{row}")] = formula
    for review_row, first, latest in ((5, 5, 6), (6, 7, 8)):
        for col, formula in {
            "B": f"=History!L{latest}-History!L{first}",
            "C": f"=History!M{latest}-History!M{first}",
            "D": f"=History!N{latest}", "E": f"=History!N{first}",
            "F": f"=D{review_row}-E{review_row}",
            "G": f"=History!O{latest}-History!O{first}",
        }.items():
            formulas[("Year Delta", f"{col}{review_row}")] = formula
        for col, formula in {
            "D": f"=History!G{latest}+B{review_row}",
            "E": f"=History!I{latest}+C{review_row}",
            "F": f"=(History!F{latest}+D{review_row})/2",
            "G": f"=(History!H{latest}+E{review_row})/2",
            "H": f"=F{review_row}/History!E{latest}*History!C{latest}",
            "I": f"=G{review_row}/History!E{latest}*History!C{latest}",
            "J": f"=D{review_row}-E{review_row}",
            "K": f"=J{review_row}-History!N{latest}",
        }.items():
            formulas[("Scenario", f"{col}{review_row}")] = formula
        for col, formula in {
            "B": f"=History!L{latest}", "C": f"=History!M{latest}",
            "D": f"=History!N{latest}", "E": f"='Year Delta'!B{review_row}",
            "F": f"='Year Delta'!C{review_row}", "G": f"=Scenario!K{review_row}",
            "H": f"=Scenario!H{review_row}", "I": f"=Scenario!J{review_row}",
        }.items():
            formulas[("Review", f"{col}{review_row}")] = formula
    # Nine distinct faults: source lineage, fiscal-year period, stock/flow
    # averaging, 53-week period length, delta sign, stress sign and summary.
    faults[("History", "F5")] = f"='Raw SEC'!N{_find(rows, 'COST', 'Inventory', '2023')}/1000000"
    faults[("History", "G5")] = f"='Raw SEC'!N{_find(rows, 'COST', 'Inventory', '2023', prior_accession=True)}/1000000"
    faults[("History", "I6")] = f"='Raw SEC'!N{_find(rows, 'COST', 'Trade payables', '2023')}/1000000"
    faults[("History", "E7")] = f"='Raw SEC'!N{_find(rows, 'WMT', 'Revenue', '2023')}/1000000"
    faults[("History", "L8")] = "=G8/E8*C8"
    faults[("History", "M6")] = "=K6/E6*365"
    faults[("Year Delta", "F5")] = "=E5-D5"
    faults[("Scenario", "E6")] = "=History!I8-C6"
    faults[("Review", "G5")] = "=Scenario!K6"
    if any(formulas[key] == wrong for key, wrong in faults.items()):
        raise ValueError("fault_does_not_change_formula")
    return formulas, faults, days


def _heading(ws, title: str, subtitle: str, headers: list[str]) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=17, bold=True, color="183153")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Aptos", size=10, italic=True, color="52687E")
    for col, name in enumerate(headers, 1):
        c = ws.cell(4, col, name)
        c.fill = PatternFill("solid", fgColor="24486C")
        c.font = Font(name="Aptos", bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[4].height = 34
    ws.freeze_panes = "D5"
    ws.column_dimensions["A"].width = 21
    for col in "BCDEFGHIJKLMNO":
        ws.column_dimensions[col].width = 18


def workbook(rows: list[dict], formulas: dict, faults: dict, days: dict, *, faulty: bool) -> Workbook:
    wb = Workbook()
    review = wb.active
    review.title = "Review"
    history = wb.create_sheet("History")
    delta = wb.create_sheet("Year Delta")
    scenario = wb.create_sheet("Scenario")
    raw = wb.create_sheet("Raw SEC")
    _heading(review, "Retail working-capital audit", "Investigate the model, repair formula defects, preserve source data and scenario inputs.",
             ["Issuer", "FY24 DIO", "FY24 DPO", "FY24 inv less AP", "DIO change", "DPO change", "Stress cash tied", "Stress DIO", "Stress net trade"])
    _heading(history, "Fiscal-year operating facts", "USD millions; source facts are from original issuer 10-Ks. Fiscal days include both endpoints.",
             ["Issuer", "Fiscal year", "Fiscal days", "Revenue", "Cost of sales", "Opening inventory", "Closing inventory", "Opening AP", "Closing AP", "Average inventory", "Average AP", "DIO", "DPO", "Inventory less AP", "Gross margin"])
    _heading(delta, "Year-over-year movement", "DIO/DPO use average balances and actual fiscal-year length; USD millions otherwise.",
             ["Issuer", "DIO change", "DPO change", "FY24 inv less AP", "FY23 inv less AP", "Net-trade change", "Gross-margin change"])
    _heading(scenario, "FY2024 stress", "Modeled shocks in USD millions, not SEC reported facts. Recalculate averages from the shocked ending balances.",
             ["Issuer", "Inventory shock", "AP shock", "Stressed ending inv", "Stressed ending AP", "Stressed avg inv", "Stressed avg AP", "Stressed DIO", "Stressed DPO", "Stressed net trade", "Incremental cash tied"])
    _heading(raw, "Frozen SEC companyfacts observations", "Authentic records only. Compare period, form, filing accession, unit and concept before linking.",
             ["Record ID", "Issuer", "Metric", "US-GAAP concept", "Period start", "Period end", "Filed", "Form", "Accession", "FY", "FP", "Frame", "Unit", "Reported USD"])
    raw.column_dimensions["D"].width = 49
    raw.column_dimensions["I"].width = 29
    for rownum, source in enumerate(rows, 5):
        for col, key in enumerate(("id", "issuer", "metric", "concept", "start", "end", "filed", "form", "accession", "fy", "fp", "frame", "unit", "value"), 1):
            raw.cell(rownum, col, source[key])
        raw[f"N{rownum}"].number_format = "#,##0;(#,##0);-"
    table = Table(displayName="RetailSECSource", ref=f"A4:N{len(rows) + 4}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    raw.add_table(table)
    for ticker, review_row, history_rows, stress in (("COST", 5, (5, 6), (1200, -400)),
                                                      ("WMT", 6, (7, 8), (5000, 2000))):
        review[f"A{review_row}"] = ticker
        delta[f"A{review_row}"] = ticker
        scenario[f"A{review_row}"] = ticker
        scenario[f"B{review_row}"] = stress[0]
        scenario[f"C{review_row}"] = stress[1]
        for year, row in zip((2023, 2024), history_rows):
            history[f"A{row}"] = ticker
            history[f"B{row}"] = year
            history[f"C{row}"] = days[("History", f"C{row}")]
    for (sheet, addr), formula in formulas.items():
        cell = wb[sheet][addr]
        cell.value = faults.get((sheet, addr), formula) if faulty else formula
        cell.font = Font(name="Aptos", size=10, color="203145")
        cell.number_format = "0.0%" if (sheet == "History" and addr[0] == "O") or (sheet == "Year Delta" and addr[0] == "G") else "#,##0.0;(#,##0.0);-"
    return wb


def build(out: Path) -> dict:
    source_rows = records()
    formulas, faults, days = plan(source_rows)
    actor, private = out / "actor", out / "private"
    actor.mkdir(parents=True, exist_ok=True)
    private.mkdir(parents=True, exist_ok=True)
    _save_deterministic(workbook(source_rows, formulas, faults, days, faulty=True), actor / "task.xlsx")
    _save_deterministic(workbook(source_rows, formulas, faults, days, faulty=False), private / "reference.xlsx")
    (private / "faults.json").write_text(json.dumps({"schema": "retail-working-capital-faults-v1",
        "fault_sites": [f"{sheet}!{addr}" for sheet, addr in faults]}, indent=2) + "\n")
    return {"source_rows": len(source_rows), "formula_cells": len(formulas), "injected_faults": len(faults), "output": str(out)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    print(json.dumps(build(parser.parse_args().output), indent=2))
