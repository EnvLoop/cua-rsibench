"""Development-only legacy SEC generator; excluded from the current rich140 protocol.

This two-issuer, one-template pool is useful for engineering controls only.
Use rich_private_replay_v1 for the current evaluator-private33-issuer corpus.
"""
import argparse,json
from pathlib import Path
from hashlib import sha256
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from .prepare_integrated_cases import packages,allocate_cases

ANNUAL=['revenue','operating_income','operating_cash_flow','capital_expenditures','cost_of_sales','assets','liabilities','equity','inventory','trade_payables','cash']
HEADERS=['Record ID','Ticker','Metric','US-GAAP concept','Start','End','Filed','Form','Accession','Fiscal year','Fiscal period','Frame','Unit','Reported USD']


def formulas(case):
    p=case['source_package'];split=case['split'];quarter=split!='train_candidate';wc=split=='final_candidate';scenario=split!='selection_candidate';out={}
    def source(metric,q=False,prior=False):
        rows=p['q3_rows'] if q else p['annual_rows'];sheet='Q3 10-Q' if q else 'Annual 10-K'
        record=p['canonical'][('q3:' if q else 'annual:')+metric]
        if prior:record=max((r for r in rows if r['metric']==metric and r['end']<p['annual_end'] and not r['start']),key=lambda r:r['end'])
        index=next(i for i,r in enumerate(rows) if r['id']==record['id'])+5
        return f"'{sheet}'!N{index}"
    def put(sheet,cell,correct,wrong=None):out[sheet+'!'+cell]={'correct':correct,'faulty':wrong or f'=({correct[1:]})*1.01'}
    for i,m in enumerate(ANNUAL):put('FY Selection',f'B{i+5}',f'={source(m)}/1000000')
    if quarter:
        for i,m in enumerate(ANNUAL[:5]):
            row=i+5;put('Q3 Selection',f'B{row}',f'={source(m,True)}/1000000')
            put('Q4 Bridge',f'B{row}',f"='FY Selection'!B{row}-'Q3 Selection'!B{row}",f"='FY Selection'!B{row}+'Q3 Selection'!B{row}")
        put('Q4 Bridge','C10','=B6/B5');put('Q4 Bridge','C11','=B7-B8','=B7+B8');put('Q4 Bridge','C12','=C11/B5')
    if wc:
        for cell,correct in [('B5',"='FY Selection'!B13"),('B6',"='FY Selection'!B14"),('B7',"='FY Selection'!B9"),('B8',"='Source Map'!B6"),
            ('B12',f'={source("inventory",prior=True)}/1000000'),('B13',f'={source("trade_payables",prior=True)}/1000000'),
            ('B9','=(B5+B12)/2/B7*B8'),('B10','=(B6+B13)/2/B7*B8'),('B11','=B5-B6')]:put('Working Capital',cell,correct)
    if scenario:
        root='Working Capital' if wc else 'FY Selection'; inv='B5' if wc else 'B13';pay='B6' if wc else 'B14'
        put('Scenario','C5',f"='{root}'!{inv}+B5",f"='{root}'!{inv}-B5")
        put('Scenario','C6',f"='{root}'!{pay}+B6",f"='{root}'!{pay}-B6")
        put('Scenario','C7',"='FY Selection'!B8*B7", "='FY Selection'!B8*(1+B7)")
        put('Scenario','C8','=C5-C6','=C5+C6');put('Scenario','C9',"='FY Selection'!B7-C7", "='FY Selection'!B7+C7")
        put('Scenario','C10',"=C9/'FY Selection'!B5")
    put('Review','B5',"='FY Selection'!B5");put('Review','B6',"='FY Selection'!B6/'FY Selection'!B5")
    put('Review','B7',"='FY Selection'!B7-'FY Selection'!B8", "='FY Selection'!B7+'FY Selection'!B8")
    put('Review','B8',"='Q4 Bridge'!B5" if quarter else "='FY Selection'!B5-'FY Selection'!B9")
    put('Review','B9',"='Q4 Bridge'!C10" if quarter else "='FY Selection'!B15")
    put('Review','B10','=Scenario!C9' if scenario else "='Q4 Bridge'!C11")
    put('Review','B11','=Scenario!C8' if scenario else "='Q4 Bridge'!C12")
    put('Audit','B5',"='FY Selection'!B10-'FY Selection'!B11-'FY Selection'!B12")
    if quarter:put('Audit','B6',"='Q4 Bridge'!B5+'Q3 Selection'!B5-'FY Selection'!B5")
    return out


def build(case, path, targets, faults):
    p=case['source_package'];split=case['split'];quarter=split!='train_candidate';wc=split=='final_candidate';scenario=split!='selection_candidate'
    names=['Review','FY Selection']+(['Q3 Selection','Q4 Bridge'] if quarter else [])+(['Working Capital'] if wc else [])+(['Scenario'] if scenario else [])+['Audit','Annual 10-K']+(['Q3 10-Q'] if quarter else [])+['Source Map']
    book=Workbook();book.remove(book.active)
    for name in names:
        sheet=book.create_sheet(name);sheet.sheet_view.showGridLines=False;sheet['A1']=name;sheet['A1'].font=Font(name='Arial',size=15,bold=True,color='173A5E')
        sheet.column_dimensions['A'].width=38;sheet.column_dimensions['B'].width=26;sheet.column_dimensions['C'].width=26;sheet.freeze_panes='A5'
        sheet['A4']='Measure';sheet['B4']='USD millions / ratio'
        for cell in sheet[4]:cell.fill=PatternFill('solid',fgColor='173A5E');cell.font=Font(name='Arial',bold=True,color='FFFFFF')
    for name,records in [('Annual 10-K',p['annual_rows'])]+([('Q3 10-Q',p['q3_rows'])] if quarter else []):
        sheet=book[name];sheet['A1']=name+' - unedited SEC filing facts';sheet['A2']=p['source_url']
        for c,value in enumerate(HEADERS,1):sheet.cell(4,c,value)
        for r,record in enumerate(records,5):
            for c,key in enumerate(('id','ticker','metric','concept','start','end','filed','form','accession','fy','fp','frame','unit','value'),1):sheet.cell(r,c,record[key])
        for c in ('D','I'):sheet.column_dimensions[c].width=45
        sheet.add_table(Table(displayName='AnnualSecFacts' if name=='Annual 10-K' else 'Q3SecFacts',ref=f'A4:N{len(records)+4}'))
    for i,m in enumerate(ANNUAL,5):book['FY Selection'].cell(i,1,m.replace('_',' '))
    if quarter:
        for i,m in enumerate(ANNUAL[:5],5):book['Q3 Selection'].cell(i,1,m.replace('_',' '));book['Q4 Bridge'].cell(i,1,m.replace('_',' '))
    labels=['FY revenue','FY operating margin','FY free cash flow','Q4 revenue bridge' if quarter else 'FY gross profit','Q4 operating margin' if quarter else 'Cash balance','Stressed free cash flow' if scenario else 'Q4 free cash flow','Stressed working capital' if scenario else 'Q4 FCF margin']
    for r,label in enumerate(labels,5):book['Review'].cell(r,1,label)
    if wc:
        labels=['Closing inventory','Closing trade payables','Cost of sales','Fiscal days','Average inventory days','Average payables days','Closing net working capital','Opening inventory','Opening payables']
        for r,label in enumerate(labels,5):book['Working Capital'].cell(r,1,label)
    if scenario:
        sheet=book['Scenario'];sheet['A1']='Modeled operating stress - not SEC facts'
        for row,key in ((5,'inventory_shock_m'),(6,'payables_shock_m'),(7,'capex_multiplier')):sheet.cell(row,1,key.replace('_',' '));sheet.cell(row,2,case['scenario'][key]);sheet.cell(row,2).fill=PatternFill('solid',fgColor='FFF0BE')
    mapping=[('CIK',f'CIK{p["cik"]:010d}'),('Fiscal days',p['fiscal_days']),('Q3 YTD days',p['q3_days']),('Annual accession',p['annual_accession']),('Annual filed',p['annual_filed']),('Q3 accession',p['q3_accession']),('Q3 filed',p['q3_filed'])]
    for r,(key,value) in enumerate(mapping,5):book['Source Map'].cell(r,1,key);book['Source Map'].cell(r,2,value)
    book['Source Map']['A13']='SEC filing facts are public. Scenarios and inserted faults are benchmark-authored.';book['Source Map']['A14']=p['source_url']
    for key,value in targets.items():
        sheet,cell=key.split('!');book[sheet][cell]=value['faulty'] if key in faults else value['correct'];book[sheet][cell].number_format='#,##0.00'
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.row!=1 and cell.row!=4:cell.font=Font(name='Arial',size=11);cell.alignment=Alignment(vertical='top')
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);book.save(path);path.chmod(0o600)


def generate(cases, output):
    output=Path(output);output.mkdir(mode=0o700,parents=True,exist_ok=False);receipts=[]
    for case in cases:
        target=formulas(case);keys=[key for key in target if not key.startswith('Audit!')];start=case['fault_offset']%len(keys)
        faults={keys[(start+i*7)%len(keys)] for i in range(case['fault_count'])}
        if len(faults)!=case['fault_count']:raise ValueError('portable_fault_pool_collision')
        root=output/case['split']/case['case_id'];root.mkdir(mode=0o700,parents=True)
        for file,selected in (('actor.xlsx',faults),('reference.xlsx',set())):build(case,root/file,target,selected)
        (root/'private-oracle.json').write_text(json.dumps({'case_id':case['case_id'],'split':case['split'],'source_group':case['source_group'],'faulted':sorted(faults),'targets':{key:value['correct'] for key,value in target.items()}})+'\n')
        (root/'task.md').write_text(f'Open the supplied workbook in Microsoft Excel for the web. Repair its unmarked filing-close formula faults. Preserve all original SEC rows and unrelated workbook content. Scenarios and faults are authored. Filing: {case["source_package"]["annual_accession"]}.\n')
        receipts.append({'case_id':case['case_id'],'split':case['split'],'actor_sha256':sha256((root/'actor.xlsx').read_bytes()).hexdigest(),'reference_sha256':sha256((root/'reference.xlsx').read_bytes()).hexdigest()})
    return receipts


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',required=True);p.add_argument('--limit',type=int,default=140);args=p.parse_args()
    cases=allocate_cases(packages())[:args.limit];rows=generate(cases,args.output)
    print(json.dumps({'schema':'sec-portable-integrated-build-v1','artifact_epoch':'openpyxl-new-bytes','cases':len(rows),'native_qualified':False}))


if __name__=='__main__':main()
