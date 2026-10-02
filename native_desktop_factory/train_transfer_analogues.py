"""Stage eight three-target WDI Desktop TRAIN packages and offline OOXML controls.

No E2B guest, LibreOffice GUI, model, or official study scorer is run here.
Country allocation and all generated files stay in ignored evaluator storage.
"""

from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
from hashlib import sha256
import json
import os
from pathlib import Path
import stat
import xml.etree.ElementTree as ET
import zipfile

from . import factory as base
from . import factory_v2
from .official_saved_verifier import verify_official
from .source import EXPECTED_SHA256, YEARS, country_facts, load
from .verify import A, P, W, S, digest, docx_content, pptx_slide_shapes, xlsx_cells


SCHEMA = "cua-native-wdi-train-transfer-analogues-v1"
PUBLIC_SCHEMA = "cua-native-wdi-train-transfer-analogues-public-v1"
TEMPLATES = {w: f"train-transfer-three-target-{w}-v1" for w in base.WORKFLOWS}
EXT = {"calc-growth": ".xlsx", "calc-risk": ".xlsx",
       "impress-deck": ".pptx", "writer-brief": ".docx"}


def _require(ok: bool, reason: str) -> None:
    if not ok:
        raise ValueError(reason)


def _sha(path: Path) -> str:
    return digest(path.read_bytes())


def _owner_only(path: Path) -> bool:
    return path.is_file() and not path.is_symlink() and stat.S_IMODE(path.stat().st_mode) & 0o077 == 0


def _calc(path: Path, facts: dict, workflow: str) -> dict:
    from openpyxl import Workbook

    book = Workbook()
    evidence = book.active
    evidence.title = "WDI source"
    evidence.append([f"{facts['name']} | source ledger"])
    evidence.append([base.source_line()])
    evidence.append(["Year", *[base.LABELS[key] for key in base.INDICATORS]])
    for year in YEARS:
        evidence.append([int(year), *[facts["years"][year][key] for key in base.INDICATORS]])
    scenario = book.create_sheet("Scenario tests")
    scenario.append(["Training committee | three linked checks"])
    scenario.append(["Use original WDI observations and the simulated policy inputs."])
    scenario.append(["Signal", "Analyst draft", "Unit", "Review note"])
    scenario.append(["Draft sheet", "Three deliberate year/rule errors", "", "Correct B5:B7 only"])
    policy = book.create_sheet("Policy inputs")
    policy.append(["Simulated controls; not WDI observations"])
    policy.append(["Control", "Value", "Unit"])
    policy.append(["Reference year", 2019, "year"])
    policy.append(["Inflation watch threshold", 5.0, "%"])
    policy.append(["Unemployment watch threshold", 6.0, "%"])
    notes = book.create_sheet("Review notes")
    notes.append(["Growth is a ratio; differences of rates are percentage points."])
    notes.append([base.source_line()])
    m = factory_v2.calculations(facts)
    if workflow == "calc-growth":
        specs = [
            ("Nominal GDP year-on-year", "=('WDI source'!B9/'WDI source'!B8-1)*100",
             "=('WDI source'!B9/'WDI source'!B7-1)*100", m["gdp_yoy"], "%"),
            ("GDP/capita five-year annualized", "=(('WDI source'!C9/'WDI source'!C4)^(1/5)-1)*100",
             "=(('WDI source'!C9/'WDI source'!C4)^(1/4)-1)*100", m["gdp_pc_cagr"], "%/year"),
            ("Inflation threshold margin", "='WDI source'!D9-'Policy inputs'!B4",
             "='WDI source'!D9-'Policy inputs'!B5",
             facts["years"]["2024"]["FP.CPI.TOTL.ZG"] - 5.0, "pp"),
        ]
    else:
        specs = [
            ("Population change since 2019", "=('WDI source'!E9/'WDI source'!E4-1)*100",
             "=('WDI source'!E9/'WDI source'!E5-1)*100", m["population_growth"], "%"),
            ("Unemployment change since 2019", "='WDI source'!F9-'WDI source'!F4",
             "='WDI source'!F9-'WDI source'!F8", m["unemployment_delta"], "pp"),
            ("Per-capita growth less inflation", "=('WDI source'!C9/'WDI source'!C8-1)*100-'WDI source'!D9",
             "=('WDI source'!C9/'WDI source'!C7-1)*100-'WDI source'!D9",
             m["pc_growth_less_inflation"], "pp"),
        ]
    targets = {}
    for i, (label, formula, wrong, value, unit) in enumerate(specs, 5):
        scenario.append([label, wrong, unit, "Audit year and rule"])
        targets[f"Scenario tests!B{i}"] = {"formula": formula, "expected_value": value}
    base._style_workbook(book, f"{facts['name']} training committee check")
    book.save(path)
    return {"targets": targets, "structure_signature": {"sheets": 4, "targets": 3,
                                                              "target_sheet": "Scenario tests"}}


def _impress(path: Path, facts: dict) -> dict:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slides = [prs.slides.add_slide(prs.slide_layouts[6]) for _ in range(7)]
    y = facts["years"]
    m = factory_v2.calculations(facts)
    titles = (
        "Training review | source and decisions", "GDP source checks", "Price source checks",
        "Labor source checks", "Analyst method", "Statements to reconcile", "Attribution",
    )
    bodies = (
        "Three decision signals must be reconciled with six years of WDI data.",
        "Compare 2024 with 2023 nominal output, not a two-year interval.",
        "Subtract annual inflation rates; report percentage points.",
        "Use the 2024 labor observation, not the previous annual print.",
        "The policy is an analyst simulation and is not World Bank guidance.",
        "Correct the three amber statements and preserve the other slides.",
        base.source_line(),
    )
    for slide, title, body in zip(slides, titles, bodies):
        factory_v2._slide(slide, title, body)
    for i, year in enumerate(YEARS):
        row = y[year]
        base._add_text(slides[1], .9, 2.8 + i * .47, 11.2, .42,
                       f"{year}: GDP USD {row['NY.GDP.MKTP.CD']/1e9:,.1f} bn", 14)
        base._add_text(slides[2], .9, 2.8 + i * .47, 11.2, .42,
                       f"{year}: inflation {row['FP.CPI.TOTL.ZG']:.2f}%", 14)
        base._add_text(slides[3], .9, 2.8 + i * .47, 11.2, .42,
                       f"{year}: unemployment {row['SL.UEM.TOTL.ZS']:.2f}%", 14)
    expected = [
        f"Output check: 2024 nominal GDP grew {m['gdp_yoy']:.2f}% year on year.",
        f"Price check: 2024 inflation changed {m['inflation_delta']:+.2f} pp year on year.",
        f"Labor check: 2024 unemployment was {y['2024']['SL.UEM.TOTL.ZS']:.2f}%.",
    ]
    wrong = [
        f"Output check: 2024 nominal GDP grew {m['gdp_yoy']+1.25:.2f}% year on year.",
        f"Price check: 2024 inflation changed {m['inflation_delta']-1.5:+.2f} pp year on year.",
        f"Labor check: 2024 unemployment was {factory_v2.stale_unemployment_rate(facts):.2f}%.",
    ]
    for i, text in enumerate(wrong):
        base._add_text(slides[5], 1.0, 2.8 + i * .95, 11.0, .72,
                       text, 21, True, "B85B1C")
    prs.save(path)
    return {"targets": dict(zip(wrong, expected)), "target_slide": 6,
            "structure_signature": {"slides": 7, "targets": 3, "target_slide": 6}}


def _writer(path: Path, facts: dict) -> dict:
    from docx import Document

    d = Document()
    d.add_heading(f"{facts['name']} | training policy review", 0)
    d.add_paragraph("Reconcile all three findings with the source table and the decision rules.")
    d.add_heading("WDI observations", 1)
    table = d.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("Year", "GDP USD bn", "Inflation %", "Unemployment %")):
        cell.text = label
    for year in YEARS:
        row = facts["years"][year]
        for cell, value in zip(table.add_row().cells,
                               (year, f"{row['NY.GDP.MKTP.CD']/1e9:,.1f}",
                                f"{row['FP.CPI.TOTL.ZG']:.2f}",
                                f"{row['SL.UEM.TOTL.ZS']:.2f}")):
            cell.text = value
    d.add_heading("Simulated decision rules", 1)
    rules = d.add_table(rows=1, cols=2)
    rules.style = "Table Grid"
    for cell, label in zip(rules.rows[0].cells, ("Code", "Interpretation")):
        cell.text = label
    for key, value in (("G", "Nominal GDP growth compares 2024 with 2023"),
                       ("P", "Inflation-rate differences use percentage points"),
                       ("L", "The labor finding uses the 2024 observation")):
        cells = rules.add_row().cells
        cells[0].text, cells[1].text = key, value
    m = factory_v2.calculations(facts)
    y = facts["years"]
    expected = [
        f"Output review: nominal GDP grew {m['gdp_yoy']:.2f}% in 2024.",
        f"Price review: annual inflation changed {m['inflation_delta']:+.2f} percentage points in 2024.",
        f"Labor review: 2024 unemployment was {y['2024']['SL.UEM.TOTL.ZS']:.2f}%.",
    ]
    wrong = [
        f"Output review: nominal GDP grew {m['gdp_yoy']+1.25:.2f}% in 2024.",
        f"Price review: annual inflation changed {m['inflation_delta']-1.5:+.2f} percentage points in 2024.",
        f"Labor review: 2024 unemployment was {factory_v2.stale_unemployment_rate(facts):.2f}%.",
    ]
    d.add_heading("Unreleased findings", 1)
    for text in wrong:
        d.add_paragraph(text, style="List Number")
    d.add_heading("Attribution", 1)
    d.add_paragraph(base.source_line())
    d.save(path)
    return {"targets": dict(zip(wrong, expected)),
            "structure_signature": {"tables": 2, "source_years": 6, "targets": 3}}


def _rewrite_zip(raw: bytes, replacements: dict[str, bytes]) -> bytes:
    import io
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, "w") as dest:
        _require(source.testzip() is None and set(replacements) <= set(source.namelist()),
                 "invalid OOXML source or replacement member")
        for name in source.namelist():
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dest.writestr(info, replacements.get(name, source.read(name)))
    return output.getvalue()


def _patch_calc(raw: bytes, oracle: dict, *, near: bool = False, unrelated: bool = False) -> bytes:
    import io
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        target_name = "xl/worksheets/sheet2.xml"
        tree = ET.fromstring(z.read(target_name))
    targets = sorted(oracle["targets"].items())
    for address, rule in targets:
        _, cell_address = address.split("!", 1)
        if near and address == targets[0][0]:
            continue
        cell = tree.find(f".//{{{S}}}sheetData/{{{S}}}row/{{{S}}}c[@r='{cell_address}']")
        _require(cell is not None, "Calc target cell missing")
        for old in list(cell):
            if old.tag in (f"{{{S}}}f", f"{{{S}}}v"):
                cell.remove(old)
        ET.SubElement(cell, f"{{{S}}}f").text = rule["formula"].lstrip("=")
        ET.SubElement(cell, f"{{{S}}}v").text = repr(rule["expected_value"])
    replacements = {target_name: ET.tostring(tree, encoding="utf-8", xml_declaration=True)}
    if unrelated:
        source_name = "xl/worksheets/sheet1.xml"
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            source_tree = ET.fromstring(z.read(source_name))
        cell = source_tree.find(f".//{{{S}}}sheetData/{{{S}}}row/{{{S}}}c[@r='B4']")
        _require(cell is not None, "source observation missing")
        value = cell.find(f"{{{S}}}v")
        _require(value is not None and value.text is not None, "source numeric value missing")
        original = float(value.text)
        value.text = str(original + max(abs(original) * .02, 1_000_000.0))
        replacements[source_name] = ET.tostring(source_tree, encoding="utf-8", xml_declaration=True)
    return _rewrite_zip(raw, replacements)


def _patch_text(raw: bytes, oracle: dict, *, near: bool = False,
                unrelated: bool = False) -> bytes:
    import io
    workflow = oracle["workflow"]
    name = (f"ppt/slides/slide{oracle['target_slide']}.xml"
            if workflow == "impress-deck" else "word/document.xml")
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        tree = ET.fromstring(z.read(name))
    tag = f"{{{A}}}t" if workflow == "impress-deck" else f"{{{W}}}t"
    targets = sorted(oracle["targets"].items())
    for old, new in targets[1 if near else 0:]:
        matches = [node for node in tree.iter(tag) if node.text == old]
        _require(len(matches) == 1, "text target not unique in OOXML")
        matches[0].text = new
    if unrelated:
        candidates = [node for node in tree.iter(tag) if node.text and
                      node.text not in oracle["targets"].values() and
                      node.text not in oracle["targets"]]
        _require(candidates, "non-target text unavailable")
        candidates[0].text += " [unrelated mutation]"
    return _rewrite_zip(raw, {name: ET.tostring(tree, encoding="utf-8", xml_declaration=True)})


def _controls(baseline: bytes, oracle: dict, salt: str) -> tuple[dict, dict[str, bytes]]:
    patch = _patch_calc if oracle["workflow"].startswith("calc-") else _patch_text
    outputs = {"positive": patch(baseline, oracle),
               "near_miss": patch(baseline, oracle, near=True),
               "unrelated_change": patch(baseline, oracle, unrelated=True)}
    results = {kind: verify_official(baseline, raw, oracle, private_salt=salt)
               for kind, raw in outputs.items()}
    _require(results["positive"]["passed"] is True and
             results["near_miss"]["passed"] is False and
             results["unrelated_change"]["passed"] is False,
             "offline saved-artifact controls did not separate 1/0/no-regression")
    return results, outputs


def _load_boundary(mapping_path: Path, candidate_root: Path) -> tuple[dict, dict, str, str]:
    mapping_raw = mapping_path.read_bytes()
    mapping = json.loads(mapping_raw)
    base.validate_private_map(mapping)
    inventory_path = candidate_root / "candidate-inventory.json"
    inventory_raw = inventory_path.read_bytes()
    inventory = json.loads(inventory_raw)
    _require(inventory.get("schema") == "cua-native-wdi-candidate-inventory-v1" and
             inventory.get("design_revision") == "v2-distinct-structures" and
             inventory.get("source_sha256") == EXPECTED_SHA256 and
             inventory.get("counts") == base.EXPECTED_COUNTS and
             len(inventory.get("tasks", [])) == 140 and
             inventory.get("private_map_sha256") == digest(base.json_bytes(mapping)),
             "original 20/20/100 candidate inventory or split map changed")
    return mapping, inventory, digest(mapping_raw), digest(inventory_raw)


def stage(mapping_path: Path, candidate_root: Path, out: Path) -> dict:
    mapping, inventory, map_sha, inventory_sha = _load_boundary(mapping_path, candidate_root)
    _require(not out.exists(), "refusing to overwrite private staged packages")
    out.mkdir(parents=True, mode=0o700)
    os.chmod(out, 0o700)
    source = load()
    rows = []
    protected = [r for r in inventory["tasks"] if r["split"] != "train"]
    protected_sources = {tuple(r["source_groups"]) for r in protected}
    protected_ids = {r["task_id"] for r in inventory["tasks"]}
    protected_templates = {r["template_group"] for r in protected}
    for workflow in base.WORKFLOWS:
        for iso in mapping["train"][:2]:
            _require((f"wdi-country:{iso}",) not in protected_sources,
                     "supplemental train source overlaps selection/final")
            facts = country_facts(source, iso)
            task_id = f"wdi-train-transfer-v1-{iso.lower()}-{workflow}"
            _require(task_id not in protected_ids and TEMPLATES[workflow] not in protected_templates,
                     "supplemental task identity or template overlaps original plan")
            path = out / task_id
            path.mkdir(mode=0o700)
            artifact = path / (task_id + EXT[workflow])
            if workflow.startswith("calc-"):
                oracle = _calc(artifact, facts, workflow)
            elif workflow == "impress-deck":
                oracle = _impress(artifact, facts)
            else:
                oracle = _writer(artifact, facts)
            base.canonicalize_ooxml(artifact)
            artifact.chmod(0o600)
            baseline = artifact.read_bytes()
            oracle.update({"schema": "cua-native-wdi-oracle-v1", "task_id": task_id,
                           "workflow": workflow, "split": "train_transfer",
                           "country_iso3": iso, "source_sha256": EXPECTED_SHA256,
                           "input_sha256": digest(baseline), "edits": 3,
                           "design_revision": "three-target-train-transfer-v1"})
            instruction = (f"In native LibreOffice, audit the {facts['name']} training file. "
                           "Correct all three linked decision calculations/statements from the "
                           "WDI source and simulated policy. Save the same file in place. "
                           "Preserve source observations, policy inputs, citations, layout, "
                           "and unrelated objects. Use only visible screenshots, mouse and keyboard.\n"
                           + base.source_line() + "\n")
            (path / "actor_task.txt").write_text(instruction)
            (path / "oracle.json").write_bytes(base.json_bytes(oracle))
            (path / "actor_task.txt").chmod(0o600)
            (path / "oracle.json").chmod(0o600)
            controls, output_bytes = _controls(baseline, oracle, mapping["variant_salt"])
            for kind, raw in output_bytes.items():
                control_path = path / (kind + EXT[workflow])
                control_path.write_bytes(raw)
                control_path.chmod(0o600)
            row = {"task_id": task_id, "split": "train_transfer", "workflow": workflow,
                   "source_groups": [f"wdi-country:{iso}"],
                   "template_group": TEMPLATES[workflow], "instance_group": task_id,
                   "input_sha256": digest(baseline),
                   "oracle_sha256": _sha(path / "oracle.json"),
                   "actor_task_sha256": _sha(path / "actor_task.txt"),
                   "offline_control_sha256": {kind: digest(raw) for kind, raw in output_bytes.items()},
                   "offline_control_status": {kind: result["passed"] for kind, result in controls.items()},
                   "gui_qualified": False, "official_final_admitted": False}
            row["package_sha256"] = digest(base.json_bytes(row))
            (path / "package.json").write_bytes(base.json_bytes(row))
            (path / "package.json").chmod(0o600)
            rows.append(row)
    _require(len(rows) == 8 and Counter(r["workflow"] for r in rows) ==
             Counter({w: 2 for w in base.WORKFLOWS}) and
             all(len({r["source_groups"][0] for r in rows if r["workflow"] == w}) == 2
                 for w in base.WORKFLOWS), "two distinct train countries per workflow required")
    private = {"schema": SCHEMA, "status": "offline_train_only_saved_controls_not_gui",
               "original_map_sha256": map_sha, "original_inventory_sha256": inventory_sha,
               "source_sha256": EXPECTED_SHA256, "rows": rows,
               "official_final_admitted": 0, "model_results": 0}
    (out / "train-transfer-inventory.private.json").write_bytes(base.json_bytes(private))
    os.chmod(out / "train-transfer-inventory.private.json", 0o600)
    return private


def audit(mapping_path: Path, candidate_root: Path, out: Path) -> dict:
    mapping, inventory, map_sha, inventory_sha = _load_boundary(mapping_path, candidate_root)
    source = load()
    private_path = out / "train-transfer-inventory.private.json"
    _require(_owner_only(private_path), "private train-transfer inventory is not owner-only")
    private = json.loads(private_path.read_bytes())
    _require(private.get("schema") == SCHEMA and
             private.get("status") == "offline_train_only_saved_controls_not_gui" and
             private.get("original_map_sha256") == map_sha and
             private.get("original_inventory_sha256") == inventory_sha and
             private.get("source_sha256") == EXPECTED_SHA256 and
             private.get("official_final_admitted") == private.get("model_results") == 0,
             "private train-transfer inventory drifted")
    rows = private["rows"]
    _require(len(rows) == 8 and len({r["task_id"] for r in rows}) == 8 and
             Counter(r["workflow"] for r in rows) == Counter({w: 2 for w in base.WORKFLOWS}),
             "eight-case train-transfer shape changed")
    protected_sources = {tuple(r["source_groups"]) for r in inventory["tasks"] if r["split"] != "train"}
    protected_ids = {r["task_id"] for r in inventory["tasks"]}
    protected_templates = {r["template_group"] for r in inventory["tasks"]}
    for row in rows:
        iso = row["source_groups"][0].split(":", 1)[1]
        _require(iso in mapping["train"][:2] and
                 tuple(row["source_groups"]) not in protected_sources and
                 row["task_id"] not in protected_ids and
                 row["template_group"] == TEMPLATES[row["workflow"]] and
                 row["template_group"] not in protected_templates and
                 row["split"] == "train_transfer" and
                 row["gui_qualified"] is False and row["official_final_admitted"] is False,
                 "source, template, or status boundary changed")
        facts = country_facts(source, iso)
        path = out / row["task_id"]
        artifact = path / (row["task_id"] + EXT[row["workflow"]])
        _require(all(_owner_only(file) for file in path.iterdir()),
                 "private task file is missing, linked, or readable by others")
        baseline = artifact.read_bytes()
        oracle = json.loads((path / "oracle.json").read_bytes())
        _require(row["input_sha256"] == digest(baseline) == oracle["input_sha256"] and
                 row["oracle_sha256"] == _sha(path / "oracle.json") and
                 row["actor_task_sha256"] == _sha(path / "actor_task.txt") and
                 row["package_sha256"] == digest(base.json_bytes({k: v for k, v in row.items()
                                                                   if k != "package_sha256"})) and
                 json.loads((path / "package.json").read_bytes()) == row and
                 oracle["country_iso3"] == iso and len(oracle["targets"]) == 3,
                 "private package source or oracle drifted")
        # Independent source-tier check: all country facts are drawn from the pinned WDI cube.
        _require(facts["iso3"] == iso and
                 oracle["source_sha256"] == EXPECTED_SHA256 and
                 source["source_sha256"] == EXPECTED_SHA256,
                 "WDI source binding changed")
        controls, output_bytes = _controls(baseline, oracle, mapping["variant_salt"])
        _require(row["offline_control_status"] == {k: v["passed"] for k, v in controls.items()} and
                 row["offline_control_sha256"] == {k: digest(v) for k, v in output_bytes.items()} and
                 all((path / (k + EXT[row["workflow"]])).read_bytes() == v
                     for k, v in output_bytes.items()),
                 "offline saved-artifact controls changed")
    return {"schema": PUBLIC_SCHEMA,
            "status": "eight_offline_three_target_train_packages_source_audited",
            "original_map_sha256": map_sha,
            "original_inventory_sha256": inventory_sha,
            "private_staged_sha256": _sha(private_path),
            "builder_source_sha256": _sha(Path(__file__)),
            "wdi_snapshot_sha256": EXPECTED_SHA256,
            "case_count": 8, "per_workflow": dict(sorted(Counter(r["workflow"] for r in rows).items())),
            "train_source_country_count": len({tuple(r["source_groups"]) for r in rows}),
            "positive_saved_artifact_pass_count": 8,
            "near_miss_rejected_count": 8,
            "unrelated_change_rejected_count": 8,
            "original_train_selection_final_counts_unchanged": [20, 20, 100],
            "source_tier": "pinned_World_Bank_WDI_observations; authored_policy_and_artifacts",
            "gui_qualified_count": 0, "sft_episode_count": 0,
            "official_final_admitted": 0, "model_results": 0}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("stage", "audit"))
    parser.add_argument("--private-map", type=Path, required=True)
    parser.add_argument("--candidate-root", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--public-out", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "stage":
        stage(args.private_map, args.candidate_root, args.out)
    public = audit(args.private_map, args.candidate_root, args.out)
    _require(not args.public_out.exists(), "refusing to overwrite public evidence")
    args.public_out.parent.mkdir(parents=True, exist_ok=True)
    args.public_out.write_bytes(base.json_bytes(public))
    print(json.dumps({"status": public["status"], "case_count": public["case_count"],
                      "gui_qualified_count": 0}, sort_keys=True))


if __name__ == "__main__":
    main()
