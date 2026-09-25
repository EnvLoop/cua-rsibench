"""Generate original, real-data native LibreOffice task packages.

Generation is offline and produces *candidates*, never runtime admission.
The private split map and all candidate files live in ignored work/.  Final
task instructions and oracle values must not be published before scoring.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import secrets
import zipfile
from datetime import datetime
from pathlib import Path

if __package__:
    from .source import (CATALOG_URL, COUNTRIES, EXPECTED_SHA256, INDICATORS,
                         LICENSE_URL, SOURCE_URL, YEARS, country_facts, load)
else:  # `python native_desktop_factory/factory.py` from the repository root.
    from source import (CATALOG_URL, COUNTRIES, EXPECTED_SHA256, INDICATORS,
                        LICENSE_URL, SOURCE_URL, YEARS, country_facts, load)


WORKFLOWS = ("calc-growth", "calc-risk", "impress-deck", "writer-brief")
EXPECTED_COUNTS = {"train": 20, "selection": 20, "final_candidate": 100}
EDIT_COUNTS = {"train": 1, "selection": 2, "final_candidate": 3, "development": 1}
LABELS = {
    "NY.GDP.MKTP.CD": "GDP (current USD)",
    "NY.GDP.PCAP.CD": "GDP per capita (current USD)",
    "FP.CPI.TOTL.ZG": "Consumer price inflation (%)",
    "SP.POP.TOTL": "Population",
    "SL.UEM.TOTL.ZS": "Unemployment (%)",
}


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def json_bytes(value: object) -> bytes:
    return (json.dumps(value, sort_keys=True, indent=2) + "\n").encode()


def canonicalize_ooxml(path: Path) -> None:
    """Normalize ZIP metadata so an identical source/map yields byte-identical packages."""
    with zipfile.ZipFile(path) as source:
        if source.testzip() is not None:
            raise ValueError("Generated Office archive is corrupt")
        members = {name: source.read(name) for name in source.namelist()}
    if "docProps/core.xml" in members:
        core = members["docProps/core.xml"]
        core = re.sub(
            rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
            rb"\g<1>2024-01-01T00:00:00Z\g<2>", core,
        )
        members["docProps/core.xml"] = core
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for name in sorted(members):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, members[name], compress_type=zipfile.ZIP_DEFLATED,
                            compresslevel=6)
    path.write_bytes(output.getvalue())


def private_map_new(path: Path) -> dict:
    if path.exists():
        raise ValueError("Refusing to overwrite an existing split map")
    countries = list(COUNTRIES)
    secrets.SystemRandom().shuffle(countries)
    mapping = {"schema": "cua-native-wdi-private-map-v1",
               "train": countries[:5], "selection": countries[5:10],
               "final_candidate": countries[10:],
               "variant_salt": secrets.token_hex(32)}
    validate_private_map(mapping)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(json_bytes(mapping))
    path.chmod(0o600)
    return mapping


def validate_private_map(mapping: object) -> None:
    if not isinstance(mapping, dict) or set(mapping) != {
        "schema", "train", "selection", "final_candidate", "variant_salt"
    } or mapping["schema"] != "cua-native-wdi-private-map-v1":
        raise ValueError("Invalid private split map schema")
    if not isinstance(mapping["variant_salt"], str) or len(mapping["variant_salt"]) < 32:
        raise ValueError("Missing private variant salt")
    all_countries = []
    for split, expected in (("train", 5), ("selection", 5), ("final_candidate", 25)):
        rows = mapping[split]
        if not isinstance(rows, list) or len(rows) != expected or not all(isinstance(x, str) for x in rows):
            raise ValueError(f"Expected {expected} country families in {split}")
        all_countries += rows
    if len(set(all_countries)) != 35 or set(all_countries) != set(COUNTRIES):
        raise ValueError("Country source families overlap or are missing")


def scenario_variant(mapping: dict, iso: str, workflow: str) -> int:
    message = f"{mapping['variant_salt']}:{iso}:{workflow}".encode()
    return int.from_bytes(hashlib.sha256(message).digest()[:4], "big") % 4


def derived(facts: dict) -> dict:
    y23, y24 = facts["years"]["2023"], facts["years"]["2024"]
    y22 = facts["years"]["2022"]
    return {
        "gdp_growth_pct": (y24["NY.GDP.MKTP.CD"] / y23["NY.GDP.MKTP.CD"] - 1) * 100,
        "gdp_pc_growth_pct": (y24["NY.GDP.PCAP.CD"] / y23["NY.GDP.PCAP.CD"] - 1) * 100,
        "population_growth_pct": (y24["SP.POP.TOTL"] / y23["SP.POP.TOTL"] - 1) * 100,
        "inflation_2024_pct": y24["FP.CPI.TOTL.ZG"],
        "inflation_change_pp": y24["FP.CPI.TOTL.ZG"] - y23["FP.CPI.TOTL.ZG"],
        "unemployment_2024_pct": y24["SL.UEM.TOTL.ZS"],
        "inflation_mean_2022_2024_pct": (
            y22["FP.CPI.TOTL.ZG"] + y23["FP.CPI.TOTL.ZG"] + y24["FP.CPI.TOTL.ZG"]
        ) / 3,
    }


def source_line() -> str:
    return ("Source: World Bank, World Development Indicators, API snapshot "
            "2026-09-25, 2019–2024; CC BY 4.0. Derived calculations are "
            "EnvLoop's, not World Bank conclusions.")


def _style_workbook(workbook, title: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    blue = PatternFill("solid", fgColor="14283B")
    for sheet in workbook:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B4"
        for col, width in {"A": 32, "B": 25, "C": 27, "D": 27,
                           "E": 25, "F": 24, "G": 31}.items():
            sheet.column_dimensions[col].width = width
        for cell in sheet[1]:
            cell.fill = blue
            cell.font = Font(name="Liberation Sans", color="FFFFFF", bold=True, size=13)
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 28
    workbook.properties.title = title
    workbook.properties.created = datetime(2024, 1, 1)
    workbook.properties.modified = datetime(2024, 1, 1)


def build_calc(path: Path, facts: dict, workflow: str, edits: int, variant: int) -> dict:
    from openpyxl import Workbook

    if workflow not in ("calc-growth", "calc-risk"):
        raise ValueError("Unknown Calc workflow")
    workbook = Workbook()
    data = workbook.active
    data.title = "WDI facts"
    data.append([f"{facts['name']} | authentic WDI observations"])
    data.append([source_line()])
    data.append(["Year", *[LABELS[indicator] for indicator in INDICATORS]])
    for year in YEARS:
        data.append([int(year), *[facts["years"][year][indicator] for indicator in INDICATORS]])
    for row in data.iter_rows(min_row=4, max_row=9, min_col=2, max_col=6):
        for cell in row:
            cell.number_format = "#,##0.00"
    review = workbook.create_sheet("Review")
    review.append([f"{facts['name']} | analyst review"])
    review.append(["Repair every marked calculation using the WDI facts tab. Preserve all source values and other cells."])
    review.append(["Measure", "Working formula", "Unit", "Source years"])
    if workflow == "calc-growth":
        specs = [
            ("Nominal GDP growth, 2024 vs 2023", "=('WDI facts'!B9/'WDI facts'!B8-1)*100", "%", "2023–2024"),
            ("GDP per-capita growth, 2024 vs 2023", "=('WDI facts'!C9/'WDI facts'!C8-1)*100", "%", "2023–2024"),
            ("Population growth, 2024 vs 2023", "=('WDI facts'!E9/'WDI facts'!E8-1)*100", "%", "2023–2024"),
        ]
        derived_keys = ("gdp_growth_pct", "gdp_pc_growth_pct", "population_growth_pct")
    else:
        specs = [
            ("Inflation change, 2024 minus 2023", "='WDI facts'!D9-'WDI facts'!D8", "percentage points", "2023–2024"),
            ("Unemployment rate, 2024", "='WDI facts'!F9", "%", "2024"),
            ("Mean inflation, 2022–2024", "=AVERAGE('WDI facts'!D7:D9)", "%", "2022–2024"),
        ]
        derived_keys = ("inflation_change_pp", "unemployment_2024_pct", "inflation_mean_2022_2024_pct")
    wrongs = ("=0", "=1", "=2", "=3")
    for i, (label, formula, unit, years) in enumerate(specs):
        review.append([label, wrongs[(variant + i) % 4] if i < edits else formula, unit, years])
        review[f"B{i+4}"].number_format = "0.000"
    review.append([])
    review.append(["Check", "All source observations on the facts tab are genuine WDI values"])
    _style_workbook(workbook, f"{facts['name']} WDI analyst review")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return {"targets": {f"Review!B{i+4}": {"formula": specs[i][1],
                                            "expected_value": derived(facts)[derived_keys[i]]}
                        for i in range(edits)},
            "all_formula_specs": {f"Review!B{i+4}": specs[i][1] for i in range(3)}}


def _add_text(slide, left, top, width, height, text, size=20, bold=False, color=None):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    para = frame.paragraphs[0]
    run = para.add_run()
    run.text = text
    run.font.name = "Liberation Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return box


def build_impress(path: Path, facts: dict, edits: int, variant: int) -> dict:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    d = derived(facts)
    slides = [prs.slides.add_slide(blank) for _ in range(4)]
    _add_text(slides[0], 0.9, 0.8, 11.6, 1.2, f"{facts['name']} | Macro review", 32, True, "14283B")
    _add_text(slides[0], 0.9, 2.1, 11.2, 1.5, "Decision support, based on authentic WDI observations. Check units and years before updating signals.", 21)
    _add_text(slides[0], 0.9, 6.7, 11.2, 0.35, source_line(), 9)
    _add_text(slides[1], 0.7, 0.5, 11.8, 0.55, "Evidence table | 2022–2024", 26, True, "14283B")
    headers = ("Year", "GDP, USD bn", "GDP/capita, USD", "Inflation, %", "Unemployment, %")
    for j, head in enumerate(headers):
        _add_text(slides[1], 0.75 + 2.47*j, 1.5, 2.4, 0.55, head, 14, True)
    for i, year in enumerate(("2022", "2023", "2024")):
        row = facts["years"][year]
        values = (year, f"{row['NY.GDP.MKTP.CD']/1e9:,.1f}",
                  f"{row['NY.GDP.PCAP.CD']:,.0f}",
                  f"{row['FP.CPI.TOTL.ZG']:.2f}",
                  f"{row['SL.UEM.TOTL.ZS']:.2f}")
        for j, value in enumerate(values):
            _add_text(slides[1], 0.75 + 2.47*j, 2.25 + i*0.85, 2.4, 0.55, value, 17)
    _add_text(slides[1], 0.75, 6.7, 11.5, 0.35, source_line(), 9)
    _add_text(slides[2], 0.7, 0.5, 11.8, 0.55, "Decision signals | verify before circulation", 26, True, "14283B")
    messages = [
        f"2024 nominal GDP growth: {d['gdp_growth_pct']:.2f}%",
        f"2024 inflation: {d['inflation_2024_pct']:.2f}%",
        f"2024 unemployment: {d['unemployment_2024_pct']:.2f}%",
    ]
    placeholders = [f"SIGNAL_{i+1}_REVIEW_{variant}" for i in range(3)]
    for i, message in enumerate(messages):
        _add_text(slides[2], 1.0, 1.65+i*1.35, 10.8, 0.8,
                  placeholders[i] if i < edits else message, 25, True, "B85B1C")
    _add_text(slides[2], 0.75, 6.7, 11.5, 0.35, source_line(), 9)
    _add_text(slides[3], 0.7, 0.5, 11.8, 0.6, "Source, definitions, and caveat", 26, True, "14283B")
    _add_text(slides[3], 0.9, 1.5, 11.2, 3.0,
              "GDP growth is nominal (current USD). Inflation and unemployment are WDI annual rates. "
              "These derived comparisons are not forecasts or World Bank judgments.", 19)
    _add_text(slides[3], 0.9, 6.55, 11.5, 0.55, source_line(), 9)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return {"targets": {placeholders[i]: messages[i] for i in range(edits)},
            "target_slide": 3}


def build_writer(path: Path, facts: dict, edits: int, variant: int) -> dict:
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    document.add_heading(f"{facts['name']} | Economic monitoring brief", 0)
    document.add_paragraph("For internal review. Reconcile each finding against the source table; preserve the source observations and methodology note.")
    document.add_heading("WDI source observations", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells,
                             ("Year", "GDP, USD bn", "GDP/capita, USD", "Inflation, %", "Population", "Unemployment, %")):
        cell.text = heading
    for year in YEARS:
        row = facts["years"][year]
        values = (year, f"{row['NY.GDP.MKTP.CD']/1e9:,.1f}",
                  f"{row['NY.GDP.PCAP.CD']:,.0f}", f"{row['FP.CPI.TOTL.ZG']:.2f}",
                  f"{row['SP.POP.TOTL']:,.0f}", f"{row['SL.UEM.TOTL.ZS']:.2f}")
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    document.add_heading("Findings for decision-makers", level=1)
    d = derived(facts)
    messages = [
        f"FINDING 1: Nominal GDP grew {d['gdp_growth_pct']:.2f}% from 2023 to 2024.",
        f"FINDING 2: Inflation changed by {d['inflation_change_pp']:+.2f} percentage points from 2023 to 2024.",
        f"FINDING 3: The 2024 unemployment rate was {d['unemployment_2024_pct']:.2f}%.",
    ]
    placeholders = [f"FINDING_{i+1}_REVIEW_{variant}" for i in range(3)]
    for i, message in enumerate(messages):
        document.add_paragraph(placeholders[i] if i < edits else message, style="List Bullet")
    document.add_heading("Method and attribution", level=1)
    document.add_paragraph("The comparison uses the same-country annual WDI observations in the table. Nominal GDP is in current USD; GDP growth is not inflation-adjusted.")
    document.add_paragraph(source_line())
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return {"targets": {placeholders[i]: messages[i] for i in range(edits)}}


def actor_instruction(facts: dict, workflow: str, edits: int) -> str:
    name = facts["name"]
    shared = (f"Open the supplied {name} document in the native LibreOffice GUI. "
              "Use only screenshots, mouse, and keyboard; save the original file in place. "
              "Do not change the authentic WDI observations, source citation, other findings, "
              "non-target objects, or unrelated files. No terminal, filesystem API, or office API is available to the actor.\n")
    if workflow == "calc-growth":
        detail = (f"In Review, repair B4 through B{3+edits} as 2024-vs-2023 percentage growth "
                  "using the indicated GDP, GDP-per-capita, and population columns on WDI facts. "
                  "Keep formulas, not pasted constants, and respect the percentage units shown.\n")
    elif workflow == "calc-risk":
        detail = (f"In Review, repair B4 through B{3+edits}: inflation change in percentage points, "
                  "2024 unemployment level, and the 2022–2024 mean inflation, in that order. "
                  "Keep formulas, not pasted constants.\n")
    elif workflow == "impress-deck":
        detail = (f"On slide 3, replace the first {edits} amber SIGNAL review text box(es) "
                  "with the corresponding 2024 GDP-growth, inflation, and unemployment statements. "
                  "Calculate growth from the source table; preserve the rest of the presentation.\n")
    else:
        detail = (f"In Findings for decision-makers, replace the first {edits} REVIEW bullet(s) "
                  "with the corresponding nominal GDP growth, inflation-change, and 2024 unemployment "
                  "statements. Round to two decimals, keep each metric's unit, and preserve the source table.\n")
    return shared + detail + source_line() + "\n"


def build_package(output: Path, source: dict, split: str, iso: str,
                  workflow: str, variant: int) -> dict:
    if split not in EDIT_COUNTS or workflow not in WORKFLOWS:
        raise ValueError("Unknown split/workflow")
    facts = country_facts(source, iso)
    edits = EDIT_COUNTS[split]
    task_id = f"wdi-native-{iso.lower()}-{workflow}"
    directory = output / split / task_id
    if directory.exists():
        raise ValueError("Refusing to overwrite a task package")
    directory.mkdir(parents=True)
    ext = ".xlsx" if workflow.startswith("calc-") else ".pptx" if workflow.startswith("impress-") else ".docx"
    artifact = directory / f"{task_id}{ext}"
    if workflow.startswith("calc-"):
        oracle = build_calc(artifact, facts, workflow, edits, variant)
    elif workflow.startswith("impress-"):
        oracle = build_impress(artifact, facts, edits, variant)
    else:
        oracle = build_writer(artifact, facts, edits, variant)
    canonicalize_ooxml(artifact)
    (directory / "actor_task.txt").write_text(actor_instruction(facts, workflow, edits))
    oracle.update({"schema": "cua-native-wdi-oracle-v1", "task_id": task_id,
                   "workflow": workflow, "split": split, "country_iso3": iso,
                   "source_sha256": EXPECTED_SHA256, "input_sha256": digest(artifact.read_bytes()),
                   "edits": edits, "variant": variant})
    (directory / "oracle.json").write_bytes(json_bytes(oracle))
    receipt = {"task_id": task_id, "split": split, "workflow": workflow,
               "source_groups": [f"wdi-country:{iso}"],
               # The current draft reuses structural layouts across splits.
               # Keep that fact visible; admission must fail until task
               # templates are redesigned or a preregistered protocol permits
               # within-template transfer.
               "template_group": f"wdi-native:{workflow}",
               "instance_group": task_id,
               "input_sha256": digest(artifact.read_bytes()),
               "oracle_sha256": digest((directory / "oracle.json").read_bytes()),
               "actor_task_sha256": digest((directory / "actor_task.txt").read_bytes())}
    receipt["package_sha256"] = digest(json_bytes(receipt))
    (directory / "package.json").write_bytes(json_bytes(receipt))
    return receipt


def generate(output: Path, mapping: dict) -> dict:
    validate_private_map(mapping)
    source = load()
    if output.exists() and any(output.iterdir()):
        raise ValueError("Refusing to overwrite nonempty output")
    receipts = []
    for split in EXPECTED_COUNTS:
        for iso in mapping[split]:
            for workflow in WORKFLOWS:
                receipts.append(build_package(output, source, split, iso, workflow,
                                              scenario_variant(mapping, iso, workflow)))
    counts = {split: sum(row["split"] == split for row in receipts) for split in EXPECTED_COUNTS}
    if counts != EXPECTED_COUNTS or len({row["task_id"] for row in receipts}) != 140:
        raise ValueError("Candidate task count/identity mismatch")
    result = {"schema": "cua-native-wdi-candidate-inventory-v1",
              "source_sha256": EXPECTED_SHA256, "source_url": SOURCE_URL,
              "catalog_url": CATALOG_URL, "license_url": LICENSE_URL,
              "private_map_sha256": digest(json_bytes(mapping)), "counts": counts,
              "status": "generated_candidates_not_runtime_admitted",
              "tasks": receipts}
    (output / "candidate-inventory.json").write_bytes(json_bytes(result))
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--private-map", type=Path, required=True)
    parser.add_argument("--new-private-map", action="store_true")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.new_private_map:
        private_map_new(args.private_map)
        print(json.dumps({"private_map_created": str(args.private_map)}, sort_keys=True))
        return
    if args.output is None:
        parser.error("--output is required when generating packages")
    mapping = json.loads(args.private_map.read_bytes())
    result = generate(args.output, mapping)
    print(json.dumps({"counts": result["counts"], "status": result["status"],
                      "inventory_sha256": digest((args.output / "candidate-inventory.json").read_bytes())},
                     sort_keys=True))


if __name__ == "__main__":
    main()
