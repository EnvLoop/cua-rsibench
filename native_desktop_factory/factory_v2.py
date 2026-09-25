"""A harder WDI candidate design with genuinely different split templates.

The v1 single-edit training layout remains exposed. Selection has a two-factor
review document; final has a multi-stage decision case with three interrelated
corrections. This module generates candidate files only, with no admissions.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

if __package__:
    from . import factory as base
    from .source import EXPECTED_SHA256, YEARS, country_facts, load
else:
    import factory as base
    from source import EXPECTED_SHA256, YEARS, country_facts, load


TEMPLATES = {
    "train": {w: f"training-single-edit-{w}-v1" for w in base.WORKFLOWS},
    "selection": {"calc-growth": "selection-three-sheet-year-contrast-v2",
                  "calc-risk": "selection-three-sheet-pressure-review-v2",
                  "impress-deck": "selection-five-slide-two-factor-brief-v2",
                  "writer-brief": "selection-one-table-two-finding-note-v2"},
    "final_candidate": {"calc-growth": "final-four-sheet-escalation-case-v2",
                        "calc-risk": "final-four-sheet-resilience-case-v2",
                        "impress-deck": "final-seven-slide-committee-case-v2",
                        "writer-brief": "final-two-table-policy-case-v2"},
}


def calculations(facts: dict) -> dict:
    y = facts["years"]
    return {
        "gdp_yoy": (y["2024"]["NY.GDP.MKTP.CD"] / y["2023"]["NY.GDP.MKTP.CD"] - 1) * 100,
        "gdp_pc_cagr": ((y["2024"]["NY.GDP.PCAP.CD"] / y["2019"]["NY.GDP.PCAP.CD"]) ** (1/5) - 1) * 100,
        "inflation_delta": y["2024"]["FP.CPI.TOTL.ZG"] - y["2023"]["FP.CPI.TOTL.ZG"],
        "population_growth": (y["2024"]["SP.POP.TOTL"] / y["2019"]["SP.POP.TOTL"] - 1) * 100,
        "unemployment_delta": y["2024"]["SL.UEM.TOTL.ZS"] - y["2019"]["SL.UEM.TOTL.ZS"],
        "unemployment_yoy": y["2024"]["SL.UEM.TOTL.ZS"] - y["2023"]["SL.UEM.TOTL.ZS"],
        "pc_growth_less_inflation": (y["2024"]["NY.GDP.PCAP.CD"] / y["2023"]["NY.GDP.PCAP.CD"] - 1) * 100 - y["2024"]["FP.CPI.TOTL.ZG"],
    }


def stale_unemployment_rate(facts: dict) -> float:
    current = round(facts["years"]["2024"]["SL.UEM.TOTL.ZS"], 2)
    for year in ("2023", "2022", "2021", "2020", "2019"):
        candidate = facts["years"][year]["SL.UEM.TOTL.ZS"]
        if round(candidate, 2) != current:
            return candidate
    return current + 1.25


def calc_document(path: Path, facts: dict, split: str, workflow: str, variant: int) -> dict:
    from openpyxl import Workbook

    book = Workbook()
    data = book.active
    data.title = "Evidence ledger"
    data.append([f"{facts['name']} | WDI evidence"])
    data.append([base.source_line()])
    data.append(["Year", *[base.LABELS[key] for key in base.INDICATORS]])
    for year in YEARS:
        data.append([int(year), *[facts["years"][year][key] for key in base.INDICATORS]])
    measures = calculations(facts)
    review = book.create_sheet("Two-factor review" if split == "selection" else "Decision ledger")
    review.append([f"{facts['name']} | {'two-factor review' if split == 'selection' else 'decision case'}"])
    review.append(["Audit every flagged calculation; keep WDI observations and simulated rules unchanged."])
    review.append(["Measure", "Working calculation", "Unit", "Period / rule"])
    if split == "selection":
        review.append(["Stage", "Analyst draft", "C: correct both formulas"])
        if workflow == "calc-growth":
            specs = [("2024 GDP growth", "=('Evidence ledger'!B9/'Evidence ledger'!B8-1)*100", measures["gdp_yoy"], "%"),
                     ("Inflation acceleration", "='Evidence ledger'!D9-'Evidence ledger'!D8", measures["inflation_delta"], "pp")]
            wrong = ["=('Evidence ledger'!B9/'Evidence ledger'!B7-1)*100", "='Evidence ledger'!D8-'Evidence ledger'!D7"]
        else:
            specs = [("2019–2024 population growth", "=('Evidence ledger'!E9/'Evidence ledger'!E4-1)*100", measures["population_growth"], "%"),
                     ("2019–2024 unemployment change", "='Evidence ledger'!F9-'Evidence ledger'!F4", measures["unemployment_delta"], "pp")]
            wrong = ["=('Evidence ledger'!E9/'Evidence ledger'!E5-1)*100", "='Evidence ledger'!F9-'Evidence ledger'!F8"]
        targets = {}
        for i, (label, formula, expected, unit) in enumerate(specs):
            review.append([label, wrong[(i+variant) % 2], unit, "Check year and unit"])
            targets[f"Two-factor review!B{i+5}"] = {"formula": formula, "expected_value": expected}
        method = book.create_sheet("Method")
        method.append(["Method note | subtraction of rates yields percentage points"])
        method.append(["One prior-year denominator is required for year-over-year growth."])
        method.append([base.source_line()])
        signature = {"sheets": 3, "targets": ["B5", "B6"], "policy_sheet": False}
    else:
        review.append(["Stage", "Draft state", "The three formulas feed a simulated decision gate"])
        review.append(["Case", f"{facts['iso3']}-2024", "Derived comparisons are not WDI forecasts"])
        policy = book.create_sheet("Risk policy" if workflow == "calc-growth" else "Resilience policy")
        policy.append(["Simulated policy parameters | not WDI observations"])
        policy.append(["Parameter", "Value", "Unit", "Use"])
        policy.append(["Reference year", 2019, "year", "Five-year comparison"])
        policy.append(["Inflation watch rate", 5.0, "%", "Escalation threshold"])
        policy.append(["Unemployment watch rate", 6.0, "%", "Secondary threshold"])
        if workflow == "calc-growth":
            specs = [("2024 nominal GDP growth", "=('Evidence ledger'!B9/'Evidence ledger'!B8-1)*100", measures["gdp_yoy"], "%"),
                     ("2019–2024 GDP/capita CAGR", "=(('Evidence ledger'!C9/'Evidence ledger'!C4)^(1/5)-1)*100", measures["gdp_pc_cagr"], "% per year"),
                     ("Inflation above watch rate", "='Evidence ledger'!D9-'Risk policy'!B4", facts["years"]["2024"]["FP.CPI.TOTL.ZG"]-5.0, "pp")]
            wrong = ["=('Evidence ledger'!B9/'Evidence ledger'!B7-1)*100",
                     "=(('Evidence ledger'!C9/'Evidence ledger'!C4)^(1/4)-1)*100",
                     "='Evidence ledger'!D9-'Risk policy'!B5"]
        else:
            specs = [("2019–2024 population growth", "=('Evidence ledger'!E9/'Evidence ledger'!E4-1)*100", measures["population_growth"], "%"),
                     ("2019–2024 unemployment change", "='Evidence ledger'!F9-'Evidence ledger'!F4", measures["unemployment_delta"], "pp"),
                     ("2024 per-capita growth minus inflation", "=('Evidence ledger'!C9/'Evidence ledger'!C8-1)*100-'Evidence ledger'!D9", measures["pc_growth_less_inflation"], "pp")]
            wrong = ["=('Evidence ledger'!E9/'Evidence ledger'!E5-1)*100",
                     "='Evidence ledger'!F9-'Evidence ledger'!F8",
                     "=('Evidence ledger'!C9/'Evidence ledger'!C7-1)*100-'Evidence ledger'!D9"]
        targets = {}
        for i, (label, formula, expected, unit) in enumerate(specs):
            review.append([label, wrong[(i+variant) % 3], unit, "2024 decision case"])
            targets[f"Decision ledger!B{i+6}"] = {"formula": formula, "expected_value": expected}
        provenance = book.create_sheet("Provenance")
        provenance.append(["Attribution and analyst interpretation boundary"])
        provenance.append([base.source_line()])
        provenance.append(["GDP and GDP per capita are nominal current-USD observations."])
        provenance.append(["Policy thresholds are simulated operational inputs."])
        signature = {"sheets": 4, "targets": ["B6", "B7", "B8"], "policy_sheet": True}
    base._style_workbook(book, f"{facts['name']} {split} analyst case")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return {"targets": targets, "structure_signature": signature}


def _slide(slide, title: str, body: str) -> None:
    base._add_text(slide, 0.8, 0.55, 11.7, 0.85, title, 27, True, "14283B")
    base._add_text(slide, 0.85, 1.6, 11.4, 1.4, body, 17)
    base._add_text(slide, 0.85, 6.8, 11.4, 0.32, base.source_line(), 8)


def impress_document(path: Path, facts: dict, split: str, variant: int) -> dict:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    count = 5 if split == "selection" else 7
    slides = [prs.slides.add_slide(prs.slide_layouts[6]) for _ in range(count)]
    y, values = facts["years"], calculations(facts)
    _slide(slides[0], f"{facts['name']} | {'Two-factor screen' if split == 'selection' else 'Committee case'}",
           "Audit the decision signals against a pinned 2019–2024 WDI evidence pack.")
    _slide(slides[1], "Output evidence", "Current-USD GDP is nominal; use year-specific values and units.")
    for i, year in enumerate(("2022", "2023", "2024") if split == "selection" else YEARS):
        v = y[year]
        base._add_text(slides[1], 0.9, 3.0+i*0.52, 11.2, 0.45,
                       f"{year}  GDP USD {v['NY.GDP.MKTP.CD']/1e9:,.1f} bn  |  GDP/capita USD {v['NY.GDP.PCAP.CD']:,.0f}", 15)
    _slide(slides[2], "Price and labor evidence", "Subtract annual rates in percentage points; do not treat an old year as current.")
    for i, year in enumerate(("2022", "2023", "2024") if split == "selection" else YEARS):
        v = y[year]
        base._add_text(slides[2], 0.9, 3.0+i*0.52, 11.2, 0.45,
                       f"{year}  inflation {v['FP.CPI.TOTL.ZG']:.2f}%  |  unemployment {v['SL.UEM.TOTL.ZS']:.2f}%", 15)
    if split == "selection":
        target_slide = 4
        _slide(slides[3], "Two-factor assessment", "Correct both amber statements before release.")
        expected = [f"2024 nominal GDP growth: {values['gdp_yoy']:.2f}%",
                    f"2024 inflation acceleration: {values['inflation_delta']:+.2f} pp"]
        wrong = [f"2024 nominal GDP growth: {values['gdp_yoy']+1.25:.2f}%",
                 f"2024 inflation acceleration: {values['inflation_delta']-1.5:+.2f} pp"]
        _slide(slides[4], "Method and attribution", "GDP growth is a ratio minus one; inflation acceleration is a difference of rates.")
    else:
        target_slide = 5
        _slide(slides[3], "Validation policy", "Use 2024 versus 2023 GDP, the inflation-rate difference, and 2024 unemployment itself.")
        _slide(slides[4], "Committee decision signals", "Three amber statements came from a stale draft. Reconcile all three.")
        expected = [f"2024 nominal GDP growth: {values['gdp_yoy']:.2f}%",
                    f"2024 inflation acceleration: {values['inflation_delta']:+.2f} pp",
                    f"2024 unemployment rate: {y['2024']['SL.UEM.TOTL.ZS']:.2f}%"]
        wrong = [f"2024 nominal GDP growth: {values['gdp_yoy']+1.25:.2f}%",
                 f"2024 inflation acceleration: {values['inflation_delta']-1.5:+.2f} pp",
                 f"2024 unemployment rate: {stale_unemployment_rate(facts):.2f}%"]
        _slide(slides[5], "Action gate", "Committee action follows after verified indicators; this is not a World Bank recommendation.")
        _slide(slides[6], "Attribution", base.source_line())
    targets = {}
    for i in range(len(expected)):
        j = (i + variant) % len(expected)
        base._add_text(slides[target_slide-1], 1.0, 3.0+i*1.0, 11.0, 0.7,
                       wrong[j], 22, True, "B85B1C")
        targets[wrong[j]] = expected[j]
    if len(targets) != len(expected):
        raise ValueError("Impress target text collided")
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return {"targets": targets, "target_slide": target_slide,
            "structure_signature": {"slides": count, "targets": len(targets)}}


def writer_document(path: Path, facts: dict, split: str, variant: int) -> dict:
    from docx import Document

    doc = Document()
    doc.add_heading(f"{facts['name']} | {'Two-factor note' if split == 'selection' else 'Policy-screen casebook'}", 0)
    doc.add_paragraph("Audit the carried-forward findings against the source values. Preserve evidence and method notes.")
    years = ("2022", "2023", "2024") if split == "selection" else YEARS
    doc.add_heading("World Development Indicators | source extract", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for c, text in zip(table.rows[0].cells, ("Year", "GDP USD bn", "GDP/capita USD", "Inflation %", "Unemployment %")):
        c.text = text
    for year in years:
        v = facts["years"][year]
        for c, text in zip(table.add_row().cells,
                           (year, f"{v['NY.GDP.MKTP.CD']/1e9:,.1f}", f"{v['NY.GDP.PCAP.CD']:,.0f}",
                            f"{v['FP.CPI.TOTL.ZG']:.2f}", f"{v['SL.UEM.TOTL.ZS']:.2f}")):
            c.text = text
    if split == "final_candidate":
        doc.add_heading("Decision-rule register | simulated", level=1)
        rules = doc.add_table(rows=1, cols=2)
        rules.style = "Table Grid"
        for c, text in zip(rules.rows[0].cells, ("Rule", "Interpretation")):
            c.text = text
        for row in (("A", "Growth uses the immediately preceding year"),
                    ("B", "A difference of rates uses percentage points"),
                    ("C", "A 2024 finding must not silently use 2019")):
            for c, text in zip(rules.add_row().cells, row):
                c.text = text
    doc.add_heading("Findings requiring correction", level=1)
    values, y = calculations(facts), facts["years"]
    expected = [f"GDP finding: nominal output grew {values['gdp_yoy']:.2f}% from 2023 to 2024.",
                f"Price finding: inflation changed by {values['inflation_delta']:+.2f} percentage points from 2023 to 2024."]
    wrong = [f"GDP finding: nominal output grew {values['gdp_yoy']+1.25:.2f}% from 2023 to 2024.",
             f"Price finding: inflation changed by {values['inflation_delta']-1.5:+.2f} percentage points from 2023 to 2024."]
    if split == "final_candidate":
        expected.append(f"Labor finding: 2024 unemployment was {y['2024']['SL.UEM.TOTL.ZS']:.2f}%.")
        wrong.append(f"Labor finding: 2024 unemployment was {stale_unemployment_rate(facts):.2f}%.")
    targets = {}
    for i in range(len(expected)):
        j = (i + variant) % len(expected)
        doc.add_paragraph(wrong[j], style="List Bullet" if split == "final_candidate" else "List Number")
        targets[wrong[j]] = expected[j]
    if len(targets) != len(expected):
        raise ValueError("Writer target text collided")
    doc.add_heading("Method and attribution", level=1)
    doc.add_paragraph("GDP figures are nominal current-USD; analyst changes in annual rates are percentage-point differences.")
    doc.add_paragraph(base.source_line())
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return {"targets": targets, "structure_signature": {"tables": 1 if split == "selection" else 2,
                                                       "source_years": len(years), "targets": len(targets)}}


def build_package(output: Path, source: dict, mapping: dict, split: str, iso: str, workflow: str) -> dict:
    variant = base.scenario_variant(mapping, iso, workflow)
    if split == "train":
        row = base.build_package(output, source, split, iso, workflow, variant)
        path = output / split / row["task_id"]
        oracle = json.loads((path / "oracle.json").read_bytes())
        oracle["structure_signature"] = (
            {"sheets": 2, "targets": 1} if workflow.startswith("calc-") else
            {"slides": 4, "targets": 1} if workflow.startswith("impress-") else
            {"tables": 1, "source_years": 6, "targets": 1}
        )
        (path / "oracle.json").write_bytes(base.json_bytes(oracle))
        row["template_group"] = TEMPLATES[split][workflow]
    else:
        facts = country_facts(source, iso)
        task_id = f"wdi-native-v2-{iso.lower()}-{workflow}"
        path = output / split / task_id
        if path.exists():
            raise ValueError("Candidate package already exists")
        path.mkdir(parents=True)
        ext = ".xlsx" if workflow.startswith("calc-") else ".pptx" if workflow.startswith("impress-") else ".docx"
        artifact = path / (task_id + ext)
        if ext == ".xlsx":
            oracle = calc_document(artifact, facts, split, workflow, variant)
        elif ext == ".pptx":
            oracle = impress_document(artifact, facts, split, variant)
        else:
            oracle = writer_document(artifact, facts, split, variant)
        base.canonicalize_ooxml(artifact)
        instruction = (f"Use the native LibreOffice GUI to audit the supplied {facts['name']} file. "
                       f"Repair all {2 if split == 'selection' else 3} incorrect decision calculations/statements "
                       "using the WDI evidence, respecting years and units. Save the same file in place. "
                       "Preserve all source observations, simulated policy inputs, citations, and unrelated objects. "
                       "Use screenshots, mouse, and keyboard only; no terminal, filesystem/Office API, or oracle.\n" +
                       base.source_line() + "\n")
        (path / "actor_task.txt").write_text(instruction)
        oracle.update({"schema": "cua-native-wdi-oracle-v1", "task_id": task_id,
                       "workflow": workflow, "split": split, "country_iso3": iso,
                       "source_sha256": EXPECTED_SHA256, "input_sha256": base.digest(artifact.read_bytes()),
                       "edits": 2 if split == "selection" else 3, "variant": variant,
                       "design_revision": "v2-distinct-structures"})
        (path / "oracle.json").write_bytes(base.json_bytes(oracle))
        row = {"task_id": task_id, "split": split, "workflow": workflow,
               "source_groups": [f"wdi-country:{iso}"],
               "template_group": TEMPLATES[split][workflow], "instance_group": task_id,
               "input_sha256": base.digest(artifact.read_bytes()),
               "oracle_sha256": base.digest((path / "oracle.json").read_bytes()),
               "actor_task_sha256": base.digest((path / "actor_task.txt").read_bytes())}
    row["oracle_sha256"] = base.digest((path / "oracle.json").read_bytes())
    without_hash = {k: v for k, v in row.items() if k != "package_sha256"}
    row["package_sha256"] = base.digest(base.json_bytes(without_hash))
    (path / "package.json").write_bytes(base.json_bytes(row))
    return row


def generate(output: Path, mapping: dict) -> dict:
    base.validate_private_map(mapping)
    source = load()
    if output.exists() and any(output.iterdir()):
        raise ValueError("Refusing to overwrite nonempty output")
    rows = [build_package(output, source, mapping, split, iso, workflow)
            for split in TEMPLATES for iso in mapping[split] for workflow in base.WORKFLOWS]
    counts = {split: sum(row["split"] == split for row in rows) for split in TEMPLATES}
    if counts != base.EXPECTED_COUNTS or len({row["task_id"] for row in rows}) != 140:
        raise ValueError("Candidate inventory shape changed")
    result = {"schema": "cua-native-wdi-candidate-inventory-v1",
              "design_revision": "v2-distinct-structures", "source_sha256": EXPECTED_SHA256,
              "private_map_sha256": base.digest(base.json_bytes(mapping)),
              "counts": counts, "status": "generated_candidates_not_runtime_admitted", "tasks": rows}
    (output / "candidate-inventory.json").write_bytes(base.json_bytes(result))
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--private-map", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = generate(args.output, json.loads(args.private_map.read_bytes()))
    print(json.dumps({"counts": result["counts"], "status": result["status"],
                      "inventory_sha256": base.digest((args.output / "candidate-inventory.json").read_bytes())},
                     sort_keys=True))


if __name__ == "__main__":
    main()
