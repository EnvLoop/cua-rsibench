"""Actual container imports/browser/artifact smoke; no benchmark score claim."""
from __future__ import annotations
import importlib
from importlib.util import find_spec
import io
import json
from pathlib import Path
import tempfile
import subprocess
import sys


def current_source_bindings(root):
    """Close real runtime dependencies in isolated, non-executing processes."""
    import os
    probes={
        'gitlab':"from gitlab_world.v066_uniform_model_workers_v14 import public_binding; value=public_binding(); result={'binding_sha256':value['binding_sha256']}",
        'odoo':"from enterprise_fallback.odoo18.native_surface_workers_v14 import public_binding; value=public_binding(); result={'binding_sha256':value['binding_sha256']}",
        'magento':"from magento_catalog_factory.native_surface_workers_v9 import public_binding; value=public_binding(); result={'binding_sha256':value['binding_sha256']}",
        'desktop':"from native_desktop_factory.native_window_current_runtime import source_manifest,public_binding; value=source_manifest(); result={'schema':value['schema'],'source_files':len(value['source_sha256s']),'source_manifest_sha256':public_binding()['source_manifest_sha256']}",
        'office':"from pathlib import Path; from tools.office_current_facade_v5 import current_sources,counterpart_registry; result={'source_files':len(current_sources(Path.cwd())),'schema':counterpart_registry(Path.cwd())['schema']}",
    }
    env={**os.environ,'PYTHONPATH':os.pathsep.join(str(p) for p in
        (root,root/'src',root/'enterprise_fallback/odoo18',root/'sec_excel_factory'))}
    results={}
    for cell,code in probes.items():
        completed=subprocess.run([sys.executable,'-c',
            'import json; '+code+'; print(json.dumps(result,sort_keys=True))'],
            cwd=root,env=env,check=True,capture_output=True,text=True,timeout=60)
        results[cell]=json.loads(completed.stdout)
    for entrypoint in ('tools.office_current_neutral_v5','tools.office_current_worker_cli_v5',
            'enterprise_fallback.odoo18.twenty_task_trial_controls_v2',
            'native_desktop_factory.native_window_current_runtime'):
        subprocess.run([sys.executable,'-m',entrypoint,'--help'],cwd=root,env=env,
            check=True,capture_output=True,timeout=60)
    return results


def run():
    assert find_spec('tinker') is None, 'Tinker must remain outside the controller image'
    from tools.build_non_tinker_source_bundle_v1 import verify
    root=Path(__file__).resolve().parents[1];manifest=verify(root)
    source_bindings=current_source_bindings(root)
    modules=['cursibench.full_study_pre_campaign_v1','cursibench.full_study_results_v1',
        'cursibench.scale_action_output_v066','gitlab_world.factory','gitlab_world.verify',
        'magento_catalog_factory.verify','native_desktop_factory.factory','native_desktop_factory.verify',
        'ppt_wdi_factory.plan','ppt_wdi_factory.verify','sec_excel_factory.extract_sec']
    for name in modules:importlib.import_module(name)
    subprocess.run([sys.executable,str(root/'sec_excel_factory/build_bridge_audit.py'),'--help'],
        check=True,capture_output=True)
    from openpyxl import Workbook,load_workbook
    from pptx import Presentation
    from docx import Document
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader
    from playwright.sync_api import sync_playwright
    import e2b
    import e2b_desktop
    from harbor.agents.base import BaseAgent
    from cursibench.harbor_adapter import BrowserAgent, TinkerBrowserAgent
    assert issubclass(BrowserAgent,BaseAgent) and issubclass(TinkerBrowserAgent,BaseAgent)
    harbor_help=subprocess.run(['harbor','--help'],check=True,capture_output=True,text=True,timeout=60)
    assert 'Usage' in harbor_help.stdout
    docker_cli=subprocess.run(['docker','--version'],check=True,capture_output=True,text=True).stdout.strip()
    compose_cli=subprocess.run(['docker','compose','version'],check=True,capture_output=True,text=True).stdout.strip()
    from native_desktop_factory import verify as artifact_verifier, source as wdi
    observations=wdi.load();assert len(observations['observations'])==1050
    verified_controls=[]
    for name,extension in [('wdi-native-mex-calc-growth','xlsx'),
            ('wdi-native-mex-impress-deck-normalized','pptx'),('wdi-native-mex-writer-brief','docx')]:
        folder=root/'native_desktop_factory/dev-fixtures'/name
        baseline=(folder/(name+'.'+extension)).read_bytes()
        oracle=json.loads((folder/'oracle.json').read_bytes())
        assert artifact_verifier.verify(baseline,(folder/'controls'/('positive.'+extension)).read_bytes(),oracle)['passed']
        assert not artifact_verifier.verify(baseline,(folder/'controls'/('near-miss.'+extension)).read_bytes(),oracle)['passed']
        verified_controls.append({'development_fixture':name,'positive_verified':True,'near_miss_rejected':True})
    with tempfile.TemporaryDirectory() as folder:
        folder=Path(folder)
        wb=Workbook();wb.active['A1']='Original value';wb.active['B1']='=1+2';wb.save(folder/'sample.xlsx')
        saved=load_workbook(folder/'sample.xlsx');assert saved.active['B1'].value=='=1+2'
        deck=Presentation();deck.slides.add_slide(deck.slide_layouts[0]);deck.save(folder/'sample.pptx')
        assert len(Presentation(folder/'sample.pptx').slides)==1
        doc=Document();doc.add_paragraph('Artifact readback smoke');doc.save(folder/'sample.docx')
        assert Document(folder/'sample.docx').paragraphs[0].text=='Artifact readback smoke'
        stream=io.BytesIO();canvas=Canvas(stream);canvas.drawString(50,750,'Build smoke: no model outcomes');canvas.save()
        assert 'Build smoke' in PdfReader(io.BytesIO(stream.getvalue())).pages[0].extract_text()
        with sync_playwright() as p:
            browser=p.chromium.launch(headless=True,args=['--no-sandbox'])
            try:
                page=browser.new_page();page.set_content('<label>Owned field<input aria-label="Owned field"></label><button>Save</button>')
                page.get_by_role('textbox',name='Owned field').fill('Saved through browser')
                assert page.get_by_role('textbox').input_value()=='Saved through browser'
                png=page.screenshot();assert png.startswith(b'\x89PNG\r\n\x1a\n')
            finally:browser.close()
    return {'schema':'envloop-non-tinker-runner-smoke-v1','source_files_verified':len(manifest['files']),
        'implementation_imports':modules,'headless_browser_action_and_screenshot_verified':True,
        'current_source_bindings':source_bindings,'current_cli_help_checked_without_execution':True,
        'docker_cli':docker_cli,'compose_cli':compose_cli,'e2b_sdks_imported_without_remote_calls':True,
        'harbor_cli_and_external_agent_imported_without_execution':True,
        'pinned_wdi_source_observations':len(observations['observations']),
        'independent_public_development_artifact_controls':verified_controls,
        'xlsx_pptx_docx_pdf_roundtrip_verified':True,'tinker_installed':find_spec('tinker') is not None,
        'tinker_api_calls':0,'model_calls':0,'real_application_qualification_claimed':False,
        'benchmark_model_results_claimed':False}


if __name__=='__main__':print(json.dumps(run(),sort_keys=True))
