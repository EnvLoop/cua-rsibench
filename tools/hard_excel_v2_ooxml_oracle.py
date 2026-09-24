"""Offline OOXML calibration for SpreadsheetBench 2 financial-model candidates.

The published gold workbook is read in trusted evaluator storage. Its answers
are not embedded in this module or emitted in receipts. This is an artifact
oracle, not proof that an agent used Excel or that cached formulas recalculated.
"""

from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from excel_web_ooxml import MAIN, read_workbook, value_at

PINNED_ARCHIVE_SHA256 = '17147ef9578cd57ce76c9a719d19da7821f3e5cb0d8f776c820f699fdcdb761c'


@dataclass(frozen=True)
class Contract:
    """Keep in evaluator memory only: it contains the expected answers."""

    source: bytes
    targets: dict[tuple[str, str], tuple[str, Decimal]]
    source_sha256: str
    gold_sha256: str
    target_sheet_counts: dict[str, int]


def digest(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def workbook(raw: bytes):
    # The bounded OOXML reader rejects duplicate/unsafe members and DTDs first.
    read_workbook(raw)
    return load_workbook(io.BytesIO(raw), read_only=False, data_only=False,
                         keep_links=False)


def populated_cells(sheet):
    # Avoid rectangular iteration: the real model has large, sparse sheets.
    return {cell.coordinate: cell for cell in sheet._cells.values()
            if cell.value is not None}


def same_cell(left, right) -> bool:
    if left is None or right is None:
        return left is None and right is None
    if left.data_type == 'f' or right.data_type == 'f':
        return left.data_type == right.data_type == 'f' and left.value == right.value
    return left.data_type == right.data_type and left.value == right.value


def defined_names(book) -> tuple:
    def entries(container):
        return tuple(sorted((name, item.attr_text, item.localSheetId)
                            for name, item in container.items()))
    return (entries(book.defined_names),
            tuple((sheet.title, entries(sheet.defined_names)) for sheet in book))


def structure_errors(source, candidate) -> list[str]:
    errors = []
    if source.sheetnames != candidate.sheetnames or source.epoch != candidate.epoch:
        errors.append('sheet_identity_or_date_system_changed')
        return errors
    if defined_names(source) != defined_names(candidate):
        errors.append('defined_names_changed')
    for name in source.sheetnames:
        first, second = source[name], candidate[name]
        if first.sheet_state != second.sheet_state:
            errors.append('sheet_visibility_changed:' + name)
        if {str(r) for r in first.merged_cells.ranges} != {str(r) for r in second.merged_cells.ranges}:
            errors.append('merged_ranges_changed:' + name)
    return errors


def answer_ranges(text: str) -> dict[str, list[tuple[int, int, int, int]]]:
    """Parse the quoted, comma-separated range syntax used by this source."""
    result: dict[str, list[tuple[int, int, int, int]]] = {}
    pattern = re.compile(r"'((?:[^']|'')+)'!([A-Z]+[1-9]\d*:[A-Z]+[1-9]\d*)")
    matches = list(pattern.finditer(text))
    residue = pattern.sub('', text).replace(',', '').strip()
    if not matches or residue:
        raise ValueError('unsupported or malformed answer_position')
    for match in matches:
        sheet = match.group(1).replace("''", "'")
        result.setdefault(sheet, []).append(range_boundaries(match.group(2)))
    return result


def in_answer_ranges(sheet: str, coordinate: str, ranges: dict) -> bool:
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    column, row = coordinate_from_string(coordinate)
    col = column_index_from_string(column)
    return any(min_col <= col <= max_col and min_row <= row <= max_row
               for min_col, min_row, max_col, max_row in ranges.get(sheet, []))


def calibrate(source_raw: bytes, gold_raw: bytes,
              answer_position: str | None = None) -> Contract:
    """Derive only blank-to-formula targets; reject any other gold mutation.

    Shared formulas are expanded by openpyxl. Floating-point spelling,
    calculated caches, and style indices are intentionally not target deltas.
    """
    source, gold = workbook(source_raw), workbook(gold_raw)
    errors = structure_errors(source, gold)
    if errors:
        raise ValueError('gold structure changed: ' + ','.join(errors[:5]))
    gold_xml = read_workbook(gold_raw)
    allowed = answer_ranges(answer_position) if answer_position else None
    targets: dict[tuple[str, str], tuple[str, Decimal]] = {}
    changes = []
    for name in source.sheetnames:
        before, after = populated_cells(source[name]), populated_cells(gold[name])
        for coordinate in before.keys() | after.keys():
            left, right = before.get(coordinate), after.get(coordinate)
            if left is None and right is not None and right.data_type == 'f':
                if allowed is not None and not in_answer_ranges(name, coordinate, allowed):
                    changes.append(f'gold_target_outside_answer_range:{name}!{coordinate}')
                    continue
                cached = value_at(gold_xml, name, coordinate)
                if cached['type'] != 'number' or not cached['formula_cache_present']:
                    changes.append(f'gold_target_cache_invalid:{name}!{coordinate}')
                    continue
                targets[(name, coordinate)] = (right.value, Decimal(cached['value']))
            elif not same_cell(left, right):
                changes.append(f'gold_unrelated_semantic_change:{name}!{coordinate}')
    if changes:
        raise ValueError(';'.join(sorted(changes)[:10]) +
                         f' ({len(changes)} total)')
    if not targets:
        raise ValueError('gold contains no blank-to-formula targets')
    counts = dict(sorted(Counter(name for name, _ in targets).items()))
    return Contract(source_raw, targets, digest(source_raw), digest(gold_raw), counts)


def cache_matches(actual: dict[str, Any], expected: Decimal) -> bool:
    if actual['type'] != 'number' or not actual['formula_cache_present']:
        return False
    observed = Decimal(actual['value'])
    # Financial-model golden caches include near-zero balance-sheet residuals.
    return abs(observed - expected) <= max(Decimal('0.0001'),
                                           abs(expected) * Decimal('0.00000001'))


def verify(contract: Contract, candidate_raw: bytes) -> dict:
    source, candidate = workbook(contract.source), workbook(candidate_raw)
    errors = structure_errors(source, candidate)
    candidate_xml = read_workbook(candidate_raw)
    if source.sheetnames == candidate.sheetnames:
        for name in source.sheetnames:
            before, after = populated_cells(source[name]), populated_cells(candidate[name])
            for coordinate in before.keys() | after.keys() | {
                    coordinate for target_sheet, coordinate in contract.targets if target_sheet == name}:
                key = name, coordinate
                left, right = before.get(coordinate), after.get(coordinate)
                if key in contract.targets:
                    formula, cache = contract.targets[key]
                    if right is None or right.data_type != 'f' or right.value != formula:
                        errors.append(f'target_formula_mismatch:{name}!{coordinate}')
                    if not cache_matches(value_at(candidate_xml, name, coordinate), cache):
                        errors.append(f'target_cache_mismatch:{name}!{coordinate}')
                elif not same_cell(left, right):
                    errors.append(f'unrelated_cell_changed:{name}!{coordinate}')
    errors.sort()
    return {
        'oracle': 'hard-excel-v2-ooxml-artifact-v1',
        'artifact_pass': not errors,
        'target_formula_cells': len(contract.targets),
        'target_sheet_counts': contract.target_sheet_counts,
        'error_count': len(errors),
        'error_examples': errors[:12],
        'source_sha256': contract.source_sha256,
        'gold_sha256': contract.gold_sha256,
        'candidate_sha256': digest(candidate_raw),
        'gui_execution_verified': False,
        'native_recalculation_verified': False,
        'formatting_and_embedded_objects_verified': False,
    }


def source_task(archive: ZipFile, task_id: str) -> tuple[dict, bytes, bytes]:
    rows = json.loads(archive.read('spreadsheetbench-v2/Financial_Model/dataset.json'))
    row = next((item for item in rows if item['id'] == task_id), None)
    if row is None:
        raise ValueError('task ID not found in Financial_Model manifest')
    prefix = 'spreadsheetbench-v2/Financial_Model/'
    for key in ('spreadsheet_path', 'golden_response_path'):
        if row[key].startswith('/') or '..' in Path(row[key]).parts:
            raise ValueError('unsafe manifest workbook path')
    return (row, archive.read(prefix + row['spreadsheet_path']),
            archive.read(prefix + row['golden_response_path']))


def mutate_cell(raw: bytes, sheet: str, coordinate: str, child: str,
                replacement: str) -> bytes:
    """Make a private negative fixture by changing one worksheet cell node."""
    parsed = read_workbook(raw)
    member = parsed['sheets'][sheet]['path']
    root = ET.fromstring(parsed['members'][member])
    cells = [cell for cell in root.iter(MAIN + 'c') if cell.get('r') == coordinate]
    if len(cells) != 1:
        raise ValueError('negative fixture cell missing or duplicated')
    node = cells[0].find(MAIN + child)
    if node is None:
        raise ValueError('negative fixture child missing')
    node.text = replacement
    edited = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    stream = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as original, ZipFile(stream, 'w') as output:
        for item in original.infolist():
            output.writestr(item, edited if item.filename == member else original.read(item.filename))
    return stream.getvalue()


def self_check(contract: Contract, gold_raw: bytes) -> dict:
    """Run source-gold positive and two disjoint synthetic negatives."""
    positive = verify(contract, gold_raw)
    target_sheet, target_coordinate = sorted(contract.targets)[0]
    wrong_target = verify(contract, mutate_cell(gold_raw, target_sheet,
                                                target_coordinate, 'f', '0'))
    book = read_workbook(gold_raw)
    unrelated = None
    for sheet, row in book['sheets'].items():
        for coordinate, cell in row['cells'].items():
            if ((sheet, coordinate) not in contract.targets and
                    cell['formula'] is None and cell['type'] == 'number'):
                unrelated = sheet, coordinate, str(Decimal(cell['value']) + 1)
                break
        if unrelated:
            break
    if unrelated is None:
        raise ValueError('no unrelated numeric cell for negative fixture')
    unrelated_edit = verify(contract, mutate_cell(gold_raw, unrelated[0],
                                                  unrelated[1], 'v', unrelated[2]))
    statuses = {
        'gold_positive': positive['artifact_pass'],
        'wrong_target_rejected': (not wrong_target['artifact_pass'] and any(
            error.startswith('target_formula_mismatch:') for error in wrong_target['error_examples'])),
        'unrelated_edit_rejected': (not unrelated_edit['artifact_pass'] and any(
            error.startswith('unrelated_cell_changed:') for error in unrelated_edit['error_examples'])),
    }
    if not all(statuses.values()):
        raise AssertionError('offline oracle positive/negative calibration failed')
    return statuses


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--archive', type=Path, required=True)
    parser.add_argument('--task-id', default='01_01')
    parser.add_argument('--candidate', type=Path)
    parser.add_argument('--self-check', action='store_true')
    args = parser.parse_args()
    hasher = hashlib.sha256()
    with args.archive.open('rb') as source_file:
        for block in iter(lambda: source_file.read(1024 * 1024), b''):
            hasher.update(block)
    if hasher.hexdigest() != PINNED_ARCHIVE_SHA256:
        raise ValueError('source archive does not match pinned revision')
    with ZipFile(args.archive) as archive:
        row, source, gold = source_task(archive, args.task_id)
    contract = calibrate(source, gold, row['answer_position'])
    result = verify(contract, args.candidate.read_bytes() if args.candidate else gold)
    result['source_task_id'] = args.task_id
    if args.self_check:
        result['synthetic_calibration'] = self_check(contract, gold)
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result['artifact_pass'] else 1)


if __name__ == '__main__':
    main()
